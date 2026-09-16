# -*- coding: utf-8 -*-
"""Offline file desensitization. Independent of accident recognition / AI analysis."""

from __future__ import annotations

import os
import re

from utils.app_paths import app_path

NER_DIRNAME = "chinese_ner"
NER_WINDOW_CHARS = 384
NER_STRIDE_CHARS = 192
NAME_LABELS = ("name", "person", "per")
ADDRESS_LABELS = ("address", "addr", "location", "loc", "gpe")
DIGIT_PATTERN = re.compile(r"[0-9０-９]")
LETTER_PATTERN = re.compile(r"[A-Za-zＡ-Ｚａ-ｚ]")
# Consecutive digits only; 18 digits, or 17 digits plus X/x checksum.
IDCARD_PATTERN = re.compile(
    r"(?<![0-9０-９])(?:[0-9０-９]{18}|[0-9０-９]{17}[XxＸｘ])(?![0-9０-９])"
)
# Consecutive 11-digit numbers only, not a slice of a longer digit run.
PHONE_PATTERN = re.compile(r"(?<![0-9０-９])[0-9０-９]{11}(?![0-9０-９])")
OPTION_KEYS = ("name", "address", "idcard", "phone", "digit", "letter")
SUPPORTED_EXCEL = (".xlsx", ".xls")
SUPPORTED_WORD = (".docx",)
SUPPORTED_EXT = SUPPORTED_EXCEL + SUPPORTED_WORD

_ner_bundle = None


class DesensitizeCancelled(Exception):
    pass


def ner_model_dir():
    return app_path("downloads", NER_DIRNAME)


def ner_model_ready():
    path = ner_model_dir()
    if not os.path.isdir(path):
        return False
    has_config = os.path.isfile(os.path.join(path, "config.json"))
    has_weight = any(
        os.path.isfile(os.path.join(path, name))
        for name in (
            "pytorch_model.bin",
            "model.safetensors",
            "model.safetensors.index.json",
        )
    )
    has_vocab = any(
        os.path.isfile(os.path.join(path, name))
        for name in ("vocab.txt", "tokenizer.json", "tokenizer_config.json")
    )
    return has_config and has_weight and has_vocab


def output_path_for(src_path):
    folder, name = os.path.split(os.path.abspath(src_path))
    return os.path.join(folder, "脱敏" + name)


def anonymize_person_name(name):
    text = str(name or "")
    if not text:
        return text
    if len(text) == 1:
        return "*"
    return text[0] + "某" + ("*" * (len(text) - 2))


def _norm_label(label):
    raw = str(label or "").strip().lower()
    if raw.startswith("b-") or raw.startswith("i-") or raw.startswith("s-") or raw.startswith("e-"):
        raw = raw[2:]
    return raw


def _is_name_label(label):
    return _norm_label(label) in NAME_LABELS


def _is_address_label(label):
    return _norm_label(label) in ADDRESS_LABELS


def _raise_if_cancelled(should_cancel):
    if should_cancel and should_cancel():
        raise DesensitizeCancelled("已取消脱敏。")


def _load_ner_bundle():
    global _ner_bundle
    if _ner_bundle is not None:
        return _ner_bundle
    if not ner_model_ready():
        raise FileNotFoundError(
            "未找到中文 NER 权重。请把模型文件放到 downloads\\chinese_ner，"
            "下载地址见 README。"
        )
    os.environ.setdefault("TRANSFORMERS_OFFLINE", "1")
    os.environ.setdefault("HF_HUB_OFFLINE", "1")
    import torch
    from transformers import AutoModelForTokenClassification, AutoTokenizer

    path = ner_model_dir()
    tokenizer = AutoTokenizer.from_pretrained(path, local_files_only=True)
    model = AutoModelForTokenClassification.from_pretrained(path, local_files_only=True)
    model.eval()
    device = torch.device("cpu")
    model.to(device)
    id2label = {int(k): str(v) for k, v in model.config.id2label.items()}
    _ner_bundle = (tokenizer, model, id2label, device, torch)
    return _ner_bundle


def _decode_window(text, tokenizer, model, id2label, device, torch):
    if not text:
        return []
    tokenize_kwargs = dict(
        return_tensors="pt",
        truncation=True,
        max_length=512,
        padding=False,
    )
    try:
        encoded = tokenizer(text, return_offsets_mapping=True, **tokenize_kwargs)
        offsets = encoded.pop("offset_mapping")[0].tolist()
    except TypeError:
        encoded = tokenizer(text, **tokenize_kwargs)
        offsets = []
        cursor = 0
        for token_id in encoded["input_ids"][0].tolist():
            piece = tokenizer.convert_ids_to_tokens([token_id])[0]
            if piece in tokenizer.all_special_tokens:
                offsets.append((0, 0))
                continue
            piece = piece.replace("##", "")
            offsets.append((cursor, cursor + len(piece)))
            cursor += len(piece)
    encoded = {
        key: value.to(device)
        for key, value in encoded.items()
        if key in ("input_ids", "attention_mask", "token_type_ids")
    }
    with torch.no_grad():
        logits = model(**encoded).logits[0]
    pred_ids = logits.argmax(dim=-1).tolist()
    spans = []
    cur_start = None
    cur_end = None
    cur_label = None
    cur_score = 0.0
    for pred_id, (start, end) in zip(pred_ids, offsets):
        if end <= start:
            continue
        label = id2label.get(int(pred_id), "O")
        kind = _norm_label(label)
        is_entity = _is_name_label(kind) or _is_address_label(kind)
        score = 1.0
        if is_entity:
            raw = str(label)
            begins = raw[:2].lower() in ("b-", "s-") or cur_label != kind
            if cur_label == kind and not begins:
                cur_end = end
            else:
                if cur_label is not None and cur_start is not None:
                    spans.append((cur_start, cur_end, cur_label, cur_score))
                cur_start, cur_end, cur_label, cur_score = start, end, kind, score
        else:
            if cur_label is not None and cur_start is not None:
                spans.append((cur_start, cur_end, cur_label, cur_score))
            cur_start = cur_end = cur_label = None
    if cur_label is not None and cur_start is not None:
        spans.append((cur_start, cur_end, cur_label, cur_score))
    return spans


def _merge_spans(spans):
    if not spans:
        return []
    spans = sorted(spans, key=lambda item: (item[0], -(item[1] - item[0])))
    merged = []
    for start, end, label, score in spans:
        if start >= end:
            continue
        hit = None
        for index, (ms, me, mlabel, mscore) in enumerate(merged):
            if end > ms and start < me:
                hit = index
                break
        if hit is None:
            merged.append([start, end, label, score])
            continue
        ms, me, mlabel, mscore = merged[hit]
        if label == mlabel:
            merged[hit] = [min(start, ms), max(end, me), label, max(score, mscore)]
            continue
        if (end - start, score) >= (me - ms, mscore):
            merged[hit] = [start, end, label, score]
    merged.sort(key=lambda item: item[0])
    compact = []
    for start, end, label, score in merged:
        if compact and compact[-1][2] == label and start <= compact[-1][1]:
            compact[-1][1] = max(compact[-1][1], end)
            compact[-1][3] = max(compact[-1][3], score)
        else:
            compact.append([start, end, label, score])
    return [(start, end, label) for start, end, label, _score in compact]


def rolling_ner_spans(text, want_name, want_address):
    if not text or not (want_name or want_address):
        return []
    tokenizer, model, id2label, device, torch = _load_ner_bundle()
    collected = []
    length = len(text)
    start = 0
    while start < length:
        end = min(length, start + NER_WINDOW_CHARS)
        chunk = text[start:end]
        for local_start, local_end, label, score in _decode_window(
            chunk, tokenizer, model, id2label, device, torch
        ):
            collected.append((local_start + start, local_end + start, label, score))
        if end >= length:
            break
        start += NER_STRIDE_CHARS
    spans = []
    for item_start, item_end, label in _merge_spans(collected):
        if want_name and _is_name_label(label):
            spans.append((item_start, item_end, "name"))
        elif want_address and _is_address_label(label):
            spans.append((item_start, item_end, "address"))
    return spans


def _star_pattern(chars, locked, source, pattern):
    for match in pattern.finditer(source):
        start, end = match.span()
        if start >= end or any(locked[start:end]):
            continue
        for index in range(start, end):
            chars[index] = "*"
            locked[index] = True


def desensitize_text(text, options):
    if text is None:
        return text
    source = str(text)
    if source == "":
        return source
    want_name = bool(options.get("name"))
    want_address = bool(options.get("address"))
    want_idcard = bool(options.get("idcard"))
    want_phone = bool(options.get("phone"))
    want_digit = bool(options.get("digit"))
    want_letter = bool(options.get("letter"))
    chars = list(source)
    locked = [False] * len(chars)
    for start, end, label in rolling_ner_spans(source, want_name, want_address):
        start = max(0, min(len(chars), start))
        end = max(start, min(len(chars), end))
        if start >= end:
            continue
        surface = "".join(chars[start:end])
        if label == "name":
            repl = anonymize_person_name(surface)
        else:
            repl = "*" * len(surface)
        if len(repl) < (end - start):
            repl = repl + ("*" * (end - start - len(repl)))
        elif len(repl) > (end - start):
            repl = repl[: end - start]
        for offset, ch in enumerate(repl):
            chars[start + offset] = ch
            locked[start + offset] = True
    if want_idcard:
        _star_pattern(chars, locked, source, IDCARD_PATTERN)
    if want_phone:
        _star_pattern(chars, locked, source, PHONE_PATTERN)
    if want_digit or want_letter:
        for index, ch in enumerate(chars):
            if locked[index]:
                continue
            if want_digit and DIGIT_PATTERN.fullmatch(ch):
                chars[index] = "*"
            elif want_letter and LETTER_PATTERN.fullmatch(ch):
                chars[index] = "*"
    return "".join(chars)


def _cell_to_text(value):
    if value is None:
        return None
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return str(value)
    if hasattr(value, "strftime"):
        return str(value)
    text = str(value)
    return text if text else None


def _replace_paragraph_text(paragraph, new_text):
    if paragraph.text == new_text:
        return
    if not paragraph.runs:
        paragraph.add_run(new_text)
        return
    paragraph.runs[0].text = new_text
    for run in paragraph.runs[1:]:
        run.text = ""


def _iter_docx_paragraphs(document):
    seen = set()

    def walk(container):
        for paragraph in getattr(container, "paragraphs", []):
            marker = id(paragraph)
            if marker in seen:
                continue
            seen.add(marker)
            yield paragraph
        for table in getattr(container, "tables", []):
            for row in table.rows:
                for cell in row.cells:
                    yield from walk(cell)

    yield from walk(document)
    for section in document.sections:
        yield from walk(section.header)
        yield from walk(section.footer)


def _desensitize_xlsx(src_path, dest_path, options, progress_cb, should_cancel):
    from openpyxl import load_workbook

    workbook = load_workbook(src_path)
    jobs = []
    for sheet in workbook.worksheets:
        skip = set()
        for merged in sheet.merged_cells.ranges:
            min_col, min_row, max_col, max_row = merged.bounds
            for row in range(min_row, max_row + 1):
                for col in range(min_col, max_col + 1):
                    if row != min_row or col != min_col:
                        skip.add((row, col))
        for row in sheet.iter_rows():
            for cell in row:
                if (cell.row, cell.column) in skip:
                    continue
                text = _cell_to_text(cell.value)
                if text is None:
                    continue
                jobs.append((cell, text))
    total = max(len(jobs), 1)
    for index, (cell, text) in enumerate(jobs):
        _raise_if_cancelled(should_cancel)
        new_text = desensitize_text(text, options)
        if new_text != text:
            cell.value = new_text
        if progress_cb and index % 8 == 0:
            percent = 18 + int(index * 72 / total)
            progress_cb(percent, "正在脱敏表格单元格 %s / %s" % (index + 1, total))
    _raise_if_cancelled(should_cancel)
    if progress_cb:
        progress_cb(94, "正在保存脱敏表格")
    workbook.save(dest_path)
    workbook.close()


def _desensitize_xls(src_path, dest_path, options, progress_cb, should_cancel):
    import xlrd
    import xlwt

    source = xlrd.open_workbook(src_path)
    output = xlwt.Workbook(encoding="utf-8")
    jobs = []
    for sheet in source.sheets():
        target = output.add_sheet(sheet.name[:31] or "Sheet")
        for row in range(sheet.nrows):
            for col in range(sheet.ncols):
                text = _cell_to_text(sheet.cell_value(row, col))
                if text is None:
                    continue
                jobs.append((target, row, col, text))
    total = max(len(jobs), 1)
    for index, (target, row, col, text) in enumerate(jobs):
        _raise_if_cancelled(should_cancel)
        new_text = desensitize_text(text, options)
        target.write(row, col, new_text)
        if progress_cb and index % 8 == 0:
            percent = 18 + int(index * 72 / total)
            progress_cb(percent, "正在脱敏表格单元格 %s / %s" % (index + 1, total))
    _raise_if_cancelled(should_cancel)
    if progress_cb:
        progress_cb(94, "正在保存脱敏表格")
    output.save(dest_path)


def _desensitize_docx(src_path, dest_path, options, progress_cb, should_cancel):
    from docx import Document

    document = Document(src_path)
    paragraphs = list(_iter_docx_paragraphs(document))
    jobs = [p for p in paragraphs if p.text]
    total = max(len(jobs), 1)
    for index, paragraph in enumerate(jobs):
        _raise_if_cancelled(should_cancel)
        old = paragraph.text
        new_text = desensitize_text(old, options)
        _replace_paragraph_text(paragraph, new_text)
        if progress_cb and index % 4 == 0:
            percent = 18 + int(index * 72 / total)
            progress_cb(percent, "正在脱敏文档段落 %s / %s" % (index + 1, total))
    _raise_if_cancelled(should_cancel)
    if progress_cb:
        progress_cb(94, "正在保存脱敏文档")
    document.save(dest_path)


def desensitize_file(src_path, options, progress_cb=None, should_cancel=None):
    src_path = os.path.abspath(src_path)
    if not os.path.isfile(src_path):
        raise FileNotFoundError("所选文件不存在。")
    ext = os.path.splitext(src_path)[1].lower()
    if ext not in SUPPORTED_EXT:
        raise ValueError("仅支持 Excel（.xlsx / .xls）和 Word（.docx）。")
    if not any(options.get(key) for key in OPTION_KEYS):
        raise ValueError("请至少勾选一项脱敏内容。")
    if (options.get("name") or options.get("address")) and not ner_model_ready():
        raise FileNotFoundError(
            "未找到中文 NER 权重。请把模型文件放到 downloads\\chinese_ner，"
            "下载地址见 README。"
        )
    dest_path = output_path_for(src_path)
    _raise_if_cancelled(should_cancel)
    if progress_cb:
        progress_cb(6, "正在准备脱敏")
    if options.get("name") or options.get("address"):
        if progress_cb:
            progress_cb(10, "正在加载中文 NER 模型")
        _load_ner_bundle()
    if progress_cb:
        progress_cb(16, "正在读取文件")
    if ext == ".xlsx":
        _desensitize_xlsx(src_path, dest_path, options, progress_cb, should_cancel)
    elif ext == ".xls":
        _desensitize_xls(src_path, dest_path, options, progress_cb, should_cancel)
    else:
        _desensitize_docx(src_path, dest_path, options, progress_cb, should_cancel)
    if progress_cb:
        progress_cb(100, "脱敏完成")
    return dest_path
