# -*- coding: utf-8 -*-
"""Adversarial AI review of accident decision / investigation Word documents."""

from __future__ import annotations

import json
import os
from urllib import error, request

from utils.ai_settings import ai_enabled, ai_request_config, sanitize_ai_settings
from utils.knowledge_retrieve import retrieve_similar_cases
from utils.knowledge_store import (
    list_pack_summaries,
    load_playbook,
    playbook_is_populated,
)


DEFAULT_CHAT_URL = "https://api.deepseek.com/chat/completions"
DEFAULT_MODEL = "deepseek-v4-pro"
MAX_DOCUMENT_CHARS = 20000
QUERY_BRIEF_CHARS = 4000
SUPPORTED_WORD = (".docx",)
DISCLAIMER = "上述为AI意见，仅供参考。"
VERDICTS = ("部分合理", "不合理", "合理")
SECTION_SPECS = (
    ("factInvestigation", "事实调查是否清晰"),
    ("legalBasis", "法律依据是否完整"),
    ("liabilityDivision", "责任划分是否合理"),
    ("conclusion", "结论是否合理"),
)

ADVERSARIAL_LEVELS = (1, 2, 3, 4, 5)
DEFAULT_ADVERSARIAL_LEVEL = 2
ADVERSARIAL_GUIDE = {
    1: {
        "name": "很包容",
        "review": (
            "对抗程度为最低、包容性最高。默认采信文书表述，仅指出明显硬伤、自相矛盾或无法成立的结论。"
            "证据或说理基本闭环时，总体判断应倾向合理，不要为质疑而质疑。"
        ),
        "discipline": (
            "纪律审查应克制：没有具体异常迹象时不得暗示受贿。"
            "仅当文书存在无法用调查疏漏解释的重大偏袒时，才可谨慎提出廉洁风险疑问。"
        ),
    },
    2: {
        "name": "较包容",
        "review": (
            "对抗程度较低。以文书为基准，重点核验明显缺口、漏引法条和责任配对硬伤。"
            "存疑之处可以指出，但若可用正常调查局限解释，不要上升为整体不合理。"
        ),
        "discipline": (
            "纪律审查保持克制，可关注异常偏轻偏重，但必须说明这只是廉洁风险提示，不得直接认定受贿。"
        ),
    },
    3: {
        "name": "适中",
        "review": (
            "对抗程度适中。该信则信、该疑则疑：事实、依据、责任、结论都要核验，"
            "有缺口就写缺口，没有缺口不要硬找。"
        ),
        "discipline": (
            "纪律审查按中等力度：检查责任畸轻畸重、回避关键不利事实、口径明显偏向一方等是否可能与廉洁风险有关，"
            "必须写明依据不足时只能存疑，不能写成已查实受贿。"
        ),
    },
    4: {
        "name": "较严",
        "review": (
            "对抗程度较高。主动寻找调查未闭合环节、法条对应不严、责任配对不自洽、结论超出证据之处，"
            "对含糊表述要从严追问。"
        ),
        "discipline": (
            "纪律审查较严：对明显偏袒、该追不问、责任明显失衡要明确提出是否存在收受贿赂导致认定偏差的嫌疑，"
            "但仍须区分嫌疑与查实，禁止捏造具体行贿事实。"
        ),
    },
    5: {
        "name": "很严",
        "review": (
            "对抗程度最高。默认怀疑文书，要求每项关键事实、法条和责任结论都能自证。"
            "说不清、证不全、配对不严即应倾向部分合理或不合理。"
        ),
        "discipline": (
            "纪律审查最严：必须专段评估报告人是否可能因收受贿赂导致认定偏差，"
            "从偏袒方向、回避不利证据、责任畸轻畸重等角度提出怀疑；"
            "没有直接证据时也要写清这是纪律风险怀疑而非已查实结论。"
        ),
    },
}


def _clamp_level(value):
    try:
        level = int(value)
    except (TypeError, ValueError):
        level = DEFAULT_ADVERSARIAL_LEVEL
    if level not in ADVERSARIAL_GUIDE:
        return DEFAULT_ADVERSARIAL_LEVEL
    return level


def _system_prompt(discipline_review):
    text = (
        "你是一名隐藏的案件复核人员，专门审查交通事故认定书、事故决定书和事故调查报告是否站得住脚。"
        "你不是协助划责的现场专家，也不是为原报告辩护的角色。"
        "你必须按指定对抗程度核验事实是否调查清楚、法律依据是否完整、责任划分是否合理、结论是否合理，"
        "指出缺口、矛盾、套用不当或依据不足之处。"
    )
    if discipline_review:
        text += (
            "本次同时启用纪律审查视角：从纪委审查角度评估报告人是否可能因收受贿赂导致认定偏差。"
            "没有直接证据时只能写嫌疑与风险，不得捏造具体受贿情节。"
        )
    text += "只能输出严格 JSON，不得返回任何额外说明或 Markdown。"
    return text


class CaseReviewCancelled(Exception):
    pass


class CaseReviewError(ValueError):
    pass


def _raise_if_cancelled(should_cancel):
    if should_cancel and should_cancel():
        raise CaseReviewCancelled("已取消事故复核。")


def output_path_for(src_path):
    folder, name = os.path.split(os.path.abspath(src_path))
    stem, ext = os.path.splitext(name)
    candidate = os.path.join(folder, "复核" + name)
    if not os.path.exists(candidate):
        return candidate
    index = 2
    while True:
        candidate = os.path.join(folder, "复核%s(%d)%s" % (stem, index, ext))
        if not os.path.exists(candidate):
            return candidate
        index += 1


def extract_docx_text(src_path):
    from docx import Document

    document = Document(src_path)
    seen = set()
    parts = []

    def walk(container):
        for paragraph in getattr(container, "paragraphs", []) or []:
            marker = id(paragraph)
            if marker in seen:
                continue
            seen.add(marker)
            text = str(paragraph.text or "").strip()
            if text:
                parts.append(text)
        for table in getattr(container, "tables", []) or []:
            for row in table.rows:
                for cell in row.cells:
                    walk(cell)

    walk(document)
    return "\n".join(parts).strip()


def _playbook_for_attachment(playbook):
    data = dict(playbook or {})
    data.pop("version", None)
    data.pop("updatedAt", None)
    return data


def _knowledge_for_review(document_text):
    playbook = load_playbook()
    query_text = str(document_text or "")[:QUERY_BRIEF_CHARS]
    similar = retrieve_similar_cases({"userReportedAccidentBrief": query_text})
    enabled = [item for item in list_pack_summaries() if item.get("enabled")]
    index_summaries = [
        {
            "name": item.get("name") or "",
            "enabled": True,
            "cardCount": int(item.get("cardCount") or 0),
            "usedRows": int(item.get("usedRows") or 0),
            "distilledAt": item.get("distilledAt") or "",
        }
        for item in enabled
    ]
    has_playbook = playbook_is_populated(playbook)
    if not has_playbook and not similar and not index_summaries:
        return None
    return {
        "usageNote": (
            "historicalLiabilityKnowledge 来自本地口径手册与已启用案件索引。"
            "只参考规律、配对、法条习惯和认定口径；不得套用历史案件具体事实。"
            "与当前文书冲突时，以当前文书为准，但必须指出冲突。"
        ),
        "enabledPackCount": len(enabled),
        "caseIndex": index_summaries,
        "playbook": _playbook_for_attachment(playbook) if has_playbook else {},
        "similarHistoricalCases": similar,
    }


def _extract_json_content(content):
    text = str(content or "").strip()
    if text.startswith("```"):
        lines = text.splitlines()
        if lines and lines[0].startswith("```"):
            lines = lines[1:]
        if lines and lines[-1].strip() == "```":
            lines = lines[:-1]
        text = "\n".join(lines).strip()
    start = text.find("{")
    end = text.rfind("}")
    if start >= 0 and end > start:
        return text[start : end + 1]
    return text


def _request_options(ai_settings):
    safe_settings = sanitize_ai_settings(ai_settings)
    request_config = ai_request_config(safe_settings)
    payload_extras = {}
    if request_config["reasoningEffort"]:
        payload_extras["reasoning_effort"] = request_config["reasoningEffort"]
    if request_config["thinkingMode"] in {"enabled", "disabled"}:
        payload_extras["thinking"] = {"type": request_config["thinkingMode"]}
    return {
        "apiKey": safe_settings["deepseekApiKey"],
        "requestUrl": request_config["requestUrl"] or DEFAULT_CHAT_URL,
        "requestTimeoutSec": request_config["requestTimeoutSec"],
        "requestConfig": request_config,
        "payloadExtras": payload_extras,
        "model": request_config["model"] or DEFAULT_MODEL,
        "temperature": request_config["temperature"],
        "appendRequirement": str(
            request_config.get("liabilityAppendRequirement", "") or ""
        ).strip(),
    }


def _build_payload(
    document_text,
    truncated,
    knowledge,
    request_options,
    adversarial_level,
    discipline_review,
):
    level = _clamp_level(adversarial_level)
    guide = ADVERSARIAL_GUIDE[level]
    requirements = [
        "必须仅返回 JSON 对象，不能返回 Markdown 或额外说明。",
        "你是案件复核角色，必须按指定对抗程度审查，不得为原报告无条件背书。",
        "不得注入或扮演系统设置中的专家角色定位。",
        "本次对抗程度为 %d/5（%s）。%s" % (level, guide["name"], guide["review"]),
        "overallVerdict 只能是：合理、部分合理、不合理。",
        "overallLine 必须是完整的第一行结论，以“复核结果：”开头，写明是否合理。",
        "四个审查段 factInvestigation / legalBasis / liabilityDivision / conclusion "
        "必须各自给出 verdict 与 reason，reason 分段说明理由，不得留空。",
        "事实调查关注经过是否查清、证据是否闭环、有无关键事实缺口。",
        "法律依据关注引用法条是否完整、是否与事实对应、有无漏引或错引。",
        "责任划分关注配对是否自洽、主次全责是否符合规则、有无责任畸轻畸重。",
        "结论关注与前文事实、依据、责任是否一致，有无超范围或依据不足。",
        "若存在 historicalLiabilityKnowledge，必须参考口径手册与案件索引中的相似案件口径；"
        "禁止照搬历史案件具体车辆、地点、伤亡或经过。",
        "文书可能已经脱敏，出现 * 号或“某”属正常，不得据此否定文书。",
    ]
    schema = {
        "overallVerdict": "合理/部分合理/不合理",
        "overallLine": "以“复核结果：”开头的完整第一行结论",
        "factInvestigation": {"verdict": "字符串", "reason": "字符串"},
        "legalBasis": {"verdict": "字符串", "reason": "字符串"},
        "liabilityDivision": {"verdict": "字符串", "reason": "字符串"},
        "conclusion": {"verdict": "字符串", "reason": "字符串"},
    }
    task = (
        "对用户上传的交通事故决定书或事故调查报告做对抗性复核。"
        "判断事实是否调查清晰、法律依据是否完整、责任划分是否合理、结论是否合理。"
    )
    if discipline_review:
        requirements.append(
            "必须额外输出 disciplineReview：从纪委纪律审查角度评估报告人是否可能因收受贿赂导致认定偏差。"
            + guide["discipline"]
        )
        schema["disciplineReview"] = {
            "verdict": "字符串，如风险较低/存在嫌疑/高度存疑",
            "reason": "字符串，说明廉洁风险与认定偏差的怀疑理由",
        }
        task += "同时按纪律审查视角评估报告人廉洁风险。"
    append_requirement = request_options.get("appendRequirement") or ""
    if append_requirement:
        requirements.append(append_requirement)
    prompt = {
        "task": task,
        "adversarialLevel": level,
        "adversarialLabel": guide["name"],
        "disciplineReviewEnabled": bool(discipline_review),
        "requirements": requirements,
        "document": {
            "truncated": bool(truncated),
            "text": document_text,
        },
        "responseSchema": schema,
    }
    if knowledge:
        prompt["historicalLiabilityKnowledge"] = knowledge
    return {
        "model": request_options["model"],
        "temperature": request_options["temperature"],
        "messages": [
            {"role": "system", "content": _system_prompt(discipline_review)},
            {"role": "user", "content": json.dumps(prompt, ensure_ascii=False)},
        ],
    }


def _post_review(
    ai_settings,
    document_text,
    truncated,
    knowledge,
    adversarial_level,
    discipline_review,
):
    request_options = _request_options(ai_settings)
    token = str(request_options["apiKey"] or "").strip()
    if not token:
        raise CaseReviewError("未配置 DeepSeek API 密钥，无法事故复核。")
    payload = _build_payload(
        document_text,
        truncated,
        knowledge,
        request_options,
        adversarial_level,
        discipline_review,
    )
    payload.update(request_options["payloadExtras"])
    body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
    req = request.Request(
        request_options["requestUrl"],
        data=body,
        headers={
            "Authorization": "Bearer %s" % token,
            "Content-Type": "application/json; charset=utf-8",
        },
        method="POST",
    )
    try:
        with request.urlopen(req, timeout=int(request_options["requestTimeoutSec"])) as response:
            raw = json.loads(response.read().decode("utf-8"))
    except error.HTTPError as exc:
        detail = ""
        try:
            detail = exc.read().decode("utf-8", errors="replace")
        except Exception:
            detail = str(exc)
        raise CaseReviewError(
            "事故复核接口失败（HTTP %s）：%s" % (exc.code, detail[:400])
        ) from exc
    except error.URLError as exc:
        raise CaseReviewError("事故复核网络失败：%s" % exc) from exc
    try:
        content = raw["choices"][0]["message"]["content"]
    except (KeyError, IndexError, TypeError) as exc:
        raise CaseReviewError("事故复核返回格式无效。") from exc
    try:
        parsed = json.loads(_extract_json_content(content))
    except json.JSONDecodeError as exc:
        raise CaseReviewError("事故复核未返回有效 JSON。") from exc
    if not isinstance(parsed, dict):
        raise CaseReviewError("事故复核未返回 JSON 对象。")
    return parsed


def _normalize_verdict(value, fallback="不合理"):
    text = str(value or "").strip()
    for item in VERDICTS:
        if item in text:
            return item
    return fallback


def _section_from(parsed, key):
    data = parsed.get(key) if isinstance(parsed, dict) else None
    if not isinstance(data, dict):
        data = {}
    verdict = str(data.get("verdict") or "").strip() or "未给出判断"
    reason = str(data.get("reason") or "").strip() or "模型未给出该段理由。"
    return {"verdict": verdict, "reason": reason}


def normalize_review_result(parsed, discipline_review=False):
    data = parsed if isinstance(parsed, dict) else {}
    overall = _normalize_verdict(data.get("overallVerdict"), "不合理")
    overall_line = str(data.get("overallLine") or "").strip()
    if not overall_line.startswith("复核结果"):
        overall_line = "复核结果：%s。" % overall
    elif overall not in overall_line:
        overall_line = "复核结果：%s。%s" % (overall, overall_line)
    sections = {}
    for key, _title in SECTION_SPECS:
        sections[key] = _section_from(data, key)
    if discipline_review:
        sections["disciplineReview"] = _section_from(data, "disciplineReview")
    return {
        "overallVerdict": overall,
        "overallLine": overall_line,
        "disciplineReview": bool(discipline_review),
        "sections": sections,
    }


def _set_run_font(run, size_pt=12, bold=False, color=None):
    from docx.oxml.ns import qn
    from docx.shared import Pt

    run.bold = bold
    run.font.size = Pt(size_pt)
    run.font.name = "宋体"
    if color is not None:
        run.font.color.rgb = color
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn("w:eastAsia"), "宋体")


def write_review_docx(dest_path, result):
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    document = Document()
    style = document.styles["Normal"]
    style.font.name = "宋体"
    style.font.size = Pt(12)
    r_pr = style.element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn("w:eastAsia"), "宋体")

    first = document.add_paragraph()
    _set_run_font(first.add_run(result["overallLine"]), size_pt=14, bold=True)

    headings = [
        ("factInvestigation", "一、事实调查是否清晰"),
        ("legalBasis", "二、法律依据是否完整"),
        ("liabilityDivision", "三、责任划分是否合理"),
        ("conclusion", "四、结论是否合理"),
    ]
    if result.get("disciplineReview"):
        headings.append(("disciplineReview", "五、纪律审查（廉洁风险与认定偏差）"))
    for key, heading in headings:
        section = result["sections"].get(key) or {
            "verdict": "未给出判断",
            "reason": "模型未给出该段理由。",
        }
        title_p = document.add_paragraph()
        _set_run_font(title_p.add_run(heading), size_pt=13, bold=True)

        judge_p = document.add_paragraph()
        _set_run_font(judge_p.add_run("判断：%s" % section["verdict"]))

        reason_p = document.add_paragraph()
        _set_run_font(reason_p.add_run(section["reason"]))

    disclaimer = document.add_paragraph()
    _set_run_font(
        disclaimer.add_run(DISCLAIMER),
        bold=True,
        color=RGBColor(0xCC, 0x00, 0x00),
    )

    document.save(dest_path)


def review_accident_document(
    src_path,
    ai_settings,
    progress_cb=None,
    should_cancel=None,
    adversarial_level=DEFAULT_ADVERSARIAL_LEVEL,
    discipline_review=True,
):
    src_path = os.path.abspath(src_path)
    if not os.path.isfile(src_path):
        raise FileNotFoundError("所选文件不存在。")
    ext = os.path.splitext(src_path)[1].lower()
    if ext not in SUPPORTED_WORD:
        raise CaseReviewError("仅支持 Word（.docx）。")
    if not ai_enabled(ai_settings):
        raise CaseReviewError("未启用 AI，无法事故复核。")
    _raise_if_cancelled(should_cancel)
    if progress_cb:
        progress_cb(8, "正在读取文书")
    text = extract_docx_text(src_path)
    if not text:
        raise CaseReviewError("所选 Word 没有可审查的正文。")
    truncated = len(text) > MAX_DOCUMENT_CHARS
    if truncated:
        text = text[:MAX_DOCUMENT_CHARS]
    _raise_if_cancelled(should_cancel)
    if progress_cb:
        progress_cb(18, "正在准备口径手册与案件索引")
    knowledge = _knowledge_for_review(text)
    _raise_if_cancelled(should_cancel)
    if progress_cb:
        progress_cb(28, "正在调用 AI 对抗性审查")
    parsed = _post_review(
        ai_settings,
        text,
        truncated,
        knowledge,
        adversarial_level,
        discipline_review,
    )
    _raise_if_cancelled(should_cancel)
    if progress_cb:
        progress_cb(88, "正在生成复核文书")
    result = normalize_review_result(parsed, discipline_review=discipline_review)
    dest_path = output_path_for(src_path)
    write_review_docx(dest_path, result)
    if progress_cb:
        progress_cb(100, "复核完成")
    return dest_path
