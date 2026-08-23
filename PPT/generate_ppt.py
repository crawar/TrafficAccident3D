# -*- coding: utf-8 -*-
"""Generate the full TrafficAccident3D briefing PPTX."""

from __future__ import annotations

import os

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
OUT_PATH = os.path.join(HERE, "TrafficAccident3D项目介绍.pptx")

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

FONT = "Microsoft YaHei"

BG = RGBColor(0x0B, 0x12, 0x20)
BG_ALT = RGBColor(0x11, 0x18, 0x27)
CARD = RGBColor(0x1E, 0x29, 0x3B)
CARD_ALT = RGBColor(0x16, 0x21, 0x32)
TEAL = RGBColor(0x14, 0xB8, 0xA6)
TEAL_LIGHT = RGBColor(0x2D, 0xD4, 0xBF)
TEAL_DARK = RGBColor(0x0F, 0x76, 0x6E)
WHITE = RGBColor(0xF8, 0xFA, 0xFC)
MUTED = RGBColor(0x94, 0xA3, 0xB8)
SLATE = RGBColor(0x33, 0x41, 0x55)
GOLD = RGBColor(0xFB, 0xBF, 0x24)
RED = RGBColor(0xF8, 0x71, 0x71)
BLUE = RGBColor(0x78, 0xB4, 0xFF)
LINE = RGBColor(0x47, 0x55, 0x69)


def _set_run_font(run, size, color, bold=False, name=FONT):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold
    r_pr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        node = r_pr.find(qn(tag))
        if node is None:
            node = etree.SubElement(r_pr, qn(tag))
        node.set("typeface", name)


def _no_line(shape):
    shape.line.fill.background()


def _fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color


def _textbox(slide, left, top, width, height, text, size, color, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    try:
        tf._txBody.bodyPr.set("anchor", {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}.get(anchor, "t"))
    except Exception:
        pass
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    _set_run_font(run, size, color, bold)
    return box


def _add_paragraph(tf, text, size, color, bold=False, space_before=6, space_after=2, align=PP_ALIGN.LEFT, level=0):
    p = tf.add_paragraph()
    p.alignment = align
    p.level = level
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    run = p.add_run()
    run.text = text
    _set_run_font(run, size, color, bold)
    return p


class Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width = SLIDE_W
        self.prs.slide_height = SLIDE_H
        self.blank = self.prs.slide_layouts[6]
        self.slides = []

    def new(self, accent=True, footer=True, title=None, subtitle=None):
        slide = self.prs.slides.add_slide(self.blank)
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SLIDE_W, SLIDE_H)
        _fill(bg, BG)
        _no_line(bg)
        if accent:
            bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.12), SLIDE_H)
            _fill(bar, TEAL)
            _no_line(bar)
        self.slides.append({"slide": slide, "footer": footer})
        if title:
            self.title(slide, title, subtitle)
        return slide

    def section(self, number, title, caption):
        slide = self.new(accent=True, footer=False)
        _textbox(slide, Inches(0.8), Inches(2.15), Inches(11.5), Inches(0.5), f"PART  {number:02d}", 16, TEAL_LIGHT, True)
        _textbox(slide, Inches(0.8), Inches(2.7), Inches(11.5), Inches(1.2), title, 40, WHITE, True)
        _textbox(slide, Inches(0.8), Inches(4.05), Inches(11.5), Inches(1.0), caption, 18, MUTED, False)
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(3.95), Inches(2.2), Inches(0.06))
        _fill(line, TEAL)
        _no_line(line)
        return slide

    def title(self, slide, title, subtitle=None):
        _textbox(slide, Inches(0.55), Inches(0.28), Inches(12.2), Inches(0.55), title, 26, WHITE, True)
        rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.55), Inches(0.86), Inches(12.2), Inches(0.025))
        _fill(rule, TEAL_DARK)
        _no_line(rule)
        if subtitle:
            _textbox(slide, Inches(0.55), Inches(0.92), Inches(12.2), Inches(0.38), subtitle, 13, MUTED, False)

    def bullets(self, slide, items, left=0.55, top=1.4, width=12.2, height=5.4, size=16, color=WHITE):
        box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
        tf = box.text_frame
        tf.word_wrap = True
        first = True
        for item in items:
            if isinstance(item, tuple):
                text, level, *rest = item
                item_size = rest[0] if rest else (size - 1 if level else size)
                item_color = rest[1] if len(rest) > 1 else (MUTED if level else color)
                item_bold = rest[2] if len(rest) > 2 else False
            else:
                text, level, item_size, item_color, item_bold = item, 0, size, color, False
            if first:
                p = tf.paragraphs[0]
                first = False
            else:
                p = tf.add_paragraph()
            p.level = level
            p.space_before = Pt(5 if level == 0 else 2)
            p.space_after = Pt(2)
            prefix = "•  " if level == 0 else "–  "
            run = p.add_run()
            run.text = prefix + text
            _set_run_font(run, item_size, item_color, item_bold)
        return box

    def card(self, slide, left, top, width, height, title, lines, title_color=TEAL_LIGHT):
        shape = slide.shapes.add_shape(
            MSO_SHAPE.ROUNDED_RECTANGLE,
            Inches(left), Inches(top), Inches(width), Inches(height),
        )
        _fill(shape, CARD)
        _no_line(shape)
        shape.adjustments[0] = 0.08
        _textbox(slide, Inches(left + 0.16), Inches(top + 0.12), Inches(width - 0.28), Inches(0.38), title, 14, title_color, True)
        box = slide.shapes.add_textbox(Inches(left + 0.16), Inches(top + 0.5), Inches(width - 0.28), Inches(height - 0.62))
        tf = box.text_frame
        tf.word_wrap = True
        first = True
        for line in lines:
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.space_before = Pt(2)
            p.space_after = Pt(3)
            run = p.add_run()
            run.text = line
            _set_run_font(run, 12, WHITE, False)
        return shape

    def table(self, slide, left, top, width, height, headers, rows, col_widths=None):
        table_shape = slide.shapes.add_table(len(rows) + 1, len(headers), Inches(left), Inches(top), Inches(width), Inches(height))
        table = table_shape.table
        if col_widths:
            total = sum(col_widths)
            for i, w in enumerate(col_widths):
                table.columns[i].width = Inches(width * w / total)
        for i, h in enumerate(headers):
            cell = table.cell(0, i)
            cell.text = ""
            p = cell.text_frame.paragraphs[0]
            run = p.add_run()
            run.text = h
            _set_run_font(run, 12, WHITE, True)
            cell.fill.solid()
            cell.fill.fore_color.rgb = TEAL_DARK
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        for r, row in enumerate(rows, start=1):
            bg = CARD if r % 2 else CARD_ALT
            for c, val in enumerate(row):
                cell = table.cell(r, c)
                cell.text = ""
                p = cell.text_frame.paragraphs[0]
                run = p.add_run()
                run.text = str(val)
                _set_run_font(run, 11, WHITE, False)
                cell.fill.solid()
                cell.fill.fore_color.rgb = bg
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        return table_shape

    def chips(self, slide, items, left, top, chip_w, chip_h, gap=0.16, fill=TEAL_DARK, text_color=WHITE):
        x = left
        for text in items:
            shape = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                Inches(x), Inches(top), Inches(chip_w), Inches(chip_h),
            )
            _fill(shape, fill)
            _no_line(shape)
            shape.adjustments[0] = 0.4
            _textbox(
                slide, Inches(x), Inches(top), Inches(chip_w), Inches(chip_h),
                text, 12, text_color, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE,
            )
            x += chip_w + gap

    def process(self, slide, steps, y=2.0, box_h=1.15):
        n = len(steps)
        gap = 0.16
        left = 0.45
        usable = 12.4
        box_w = (usable - gap * (n - 1) - 0.22 * (n - 1)) / n
        x = left
        for i, (title, body) in enumerate(steps):
            shape = slide.shapes.add_shape(
                MSO_SHAPE.ROUNDED_RECTANGLE,
                Inches(x), Inches(y), Inches(box_w), Inches(box_h),
            )
            _fill(shape, CARD)
            _no_line(shape)
            shape.adjustments[0] = 0.1
            _textbox(slide, Inches(x + 0.1), Inches(y + 0.12), Inches(box_w - 0.18), Inches(0.32), f"{i + 1}. {title}", 13, TEAL_LIGHT, True)
            _textbox(slide, Inches(x + 0.1), Inches(y + 0.46), Inches(box_w - 0.18), Inches(box_h - 0.56), body, 11, WHITE, False)
            if i < n - 1:
                arrow = slide.shapes.add_shape(
                    MSO_SHAPE.RIGHT_ARROW,
                    Inches(x + box_w + 0.02), Inches(y + box_h / 2 - 0.1),
                    Inches(0.18), Inches(0.2),
                )
                _fill(arrow, TEAL)
                _no_line(arrow)
            x += box_w + gap + 0.22

    def finish(self, path):
        total = len(self.slides)
        for i, item in enumerate(self.slides, start=1):
            if not item["footer"]:
                continue
            slide = item["slide"]
            line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.55), Inches(7.12), Inches(12.2), Inches(0.015))
            _fill(line, SLATE)
            _no_line(line)
            _textbox(slide, Inches(0.55), Inches(7.16), Inches(8.5), Inches(0.28), "TrafficAccident3D  ·  六支队事故智脑", 10, MUTED)
            _textbox(slide, Inches(10.4), Inches(7.16), Inches(2.35), Inches(0.28), f"{i:02d}  /  {total:02d}", 10, TEAL_LIGHT, True, PP_ALIGN.RIGHT)
        self.prs.save(path)


def build():
    d = Deck()

    # 01 cover
    s = d.new(accent=False, footer=False)
    band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.22), SLIDE_H)
    _fill(band, TEAL)
    _no_line(band)
    _textbox(s, Inches(0.8), Inches(1.55), Inches(11.5), Inches(0.4), "ACCIDENT MIND  ·  AI  ·  TECHNICAL BRIEFING", 14, TEAL_LIGHT, True)
    _textbox(s, Inches(0.8), Inches(2.15), Inches(11.8), Inches(1.5), "六支队\n事故智脑", 40, WHITE, True)
    _textbox(s, Inches(0.8), Inches(3.85), Inches(11.5), Inches(0.45), "TrafficAccident3D  完整业务逻辑 · 系统设置 · 技术栈说明", 18, MUTED)
    d.chips(
        s,
        ["航拍识别", "人工校对", "运动路径", "三维 HTML", "AI 责任分析"],
        0.8, 4.55, 2.05, 0.42,
    )
    _textbox(s, Inches(0.8), Inches(6.55), Inches(11.5), Inches(0.35), "内部培训 / 技术交底  ·  不含 API 密钥  ·  可随代码同步再生成", 12, MUTED)

    # 02 agenda
    s = d.new(title="目录", subtitle="按真实代码模块组织，覆盖从启动到开源发布的全链路")
    agenda = [
        ("01  定位与价值", "给谁用、解决什么、一天怎么干完一件案子"),
        ("02  技术栈与架构", "桌面端、识别、三维、AI、本地服务如何拼在一起"),
        ("03  主流程与识别", "上传、YOLO、车辆类型推断、进入校对"),
        ("04  结果确认与编辑", "五图层、快捷键、必填校验、缓存"),
        ("05  系统设置全表", "识别 / AI / 专家 / 端口 / 动画 / 车型规格"),
        ("06  三维报告与播放", "测量、编号、责任、报告、运动圆角与车头约束"),
        ("07  数据、打包与安全", "JSON、history、密钥隔离、GitHub 与 PyInstaller"),
    ]
    for i, (t, b) in enumerate(agenda):
        col = i % 2
        row = i // 2
        d.card(s, 0.55 + col * 6.25, 1.35 + row * 1.38, 6.05, 1.28, t, [b])

    # 03 section
    d.section(1, "定位与价值", "一张航拍图，从现场识别走到可交互三维现场与专家发言。")

    s = d.new(title="项目是什么", subtitle="面向高速公路交警的「事故智脑」桌面工具")
    d.card(s, 0.55, 1.4, 4.0, 5.2, "一句话", [
        "事故智脑把人工智能接到交通事故现场。",
        "上传一张事故现场航拍图，",
        "自动检出车辆（支持旋转框），",
        "人工校对车辆、道路线、标记物、运动路径和案情后，",
        "生成可离线打开的三维 HTML。",
        "",
        "配置大模型后，可按专家角色一次生成分析与发言。",
    ])
    d.card(s, 4.75, 1.4, 4.0, 5.2, "给谁用", [
        "高速公路交警（六支队场景）。",
        "需要快速把航拍图变成可讲解、可测量、可回放的数字现场。",
        "",
        "案件侦办：还原经过与责任方向。",
        "秩序管理：提炼违法行为与整治动作。",
        "护路联防：看道路因素与设施缺口。",
        "指挥调度：看监控/GPS 能发现的动态风险。",
    ])
    d.card(s, 8.95, 1.4, 3.8, 5.2, "明确不做", [
        "不是行车记录仪视频分析。",
        "不是实时监控接入。",
        "不替代现场勘查笔录。",
        "不把 API 密钥写入仓库或安装包。",
        "GitHub 不托管超过 100MB 的 YOLO 权重。",
        "",
        "无密钥时：识别 + 校对 + 三维建模仍完整可用。",
    ])

    s = d.new(title="日常办案路径（交警视角）", subtitle="主窗口四个入口：上传图片、查看历史、建模预览、系统设置")
    d.process(
        s,
        [
            ("上传航拍", "选现场俯视图，后台线程跑 YOLO"),
            ("校对现场", "车辆与道路线必填，标记/路径/案情可选"),
            ("识别确认", "生成三维 HTML，可选一次 AI 请求"),
            ("现场讲解", "浏览器打开：测量、播放、责任、报告"),
        ],
        y=1.45,
        box_h=1.35,
    )
    d.card(s, 0.55, 3.15, 6.05, 3.4, "主窗口按钮", [
        "上传图片：选择图片后自动开始识别（需 downloads 中有 .pt）。",
        "查看历史：打开 history 下已生成的 HTML。",
        "建模预览：单车 Three.js 预览，底部切换 8 种标准车型，与导出场景同源。",
        "系统设置：六个 Tab，见后文专章。",
        "端口被占用时，主界面会提示「端口被占用请另行设置」。",
    ])
    d.card(s, 6.8, 3.15, 5.95, 3.4, "产物落在哪里", [
        "三维现场：history/YYYYMMDD_HHMMSS.html",
        "带测量线副本：history/{用户名}_{时间}(line).html",
        "校对缓存：与图片绑定的标注历史，含车辆尺寸、车道宽、案情、路径。",
        "配置：config/*.json（密钥仅本地 ai_settings.json）。",
        "专家头像：ExpertPic/（默认四张 PNG + 自定义）。",
    ])

    # 04 section tech
    d.section(2, "技术栈与架构", "桌面 GUI、视觉识别、三维导出、可选大模型，全部跑在本机。")

    s = d.new(title="技术栈总览", subtitle="运行依赖写在 requirements.txt；三维与打包工具不进入该文件")
    d.table(
        s, 0.55, 1.35, 12.2, 5.4,
        ["层次", "技术", "版本/说明", "职责"],
        [
            ["语言 / 运行", "Python + Conda", "3.10，环境在项目内 env/", "create.bat 创建，start.bat 启动"],
            ["桌面 GUI", "PySide6 (Qt)", ">= 6.7.0", "主窗、校对、系统设置、快捷键"],
            ["视觉识别", "Ultralytics YOLO11", ">= 8.2.0，权重 .pt", "检出车辆；优先 OBB 旋转框"],
            ["图像", "OpenCV + Pillow + NumPy", "cv2 / PIL / np", "读图、画框、采样车身颜色、坐标换算"],
            ["三维现场", "Three.js + GLB", "模板内嵌，离线可开", "场景、测量、运动、责任标牌"],
            ["AI（可选）", "DeepSeek Chat Completions", "默认 deepseek-v4-pro", "一次请求：结构化责任 + 专家发言"],
            ["本地服务", "Python HTTP Server", "默认端口 8765", "HTML「保存结果」写回 history"],
            ["打包", "PyInstaller", "build.bat", "release/TrafficAccident3D/"],
        ],
        col_widths=[1.6, 2.6, 3.2, 4.8],
    )

    s = d.new(title="运行时架构", subtitle="main.py 拉起 Qt 应用，同时启动测量保存服务")
    d.card(s, 0.55, 1.4, 4.0, 5.2, "进程内模块", [
        "main.py：QApplication、全局字体兜底、start/stop 测量服务。",
        "gui/main_window.py：主界面与识别线程 ProcessThread。",
        "core/yolo_detector.py：加载 .pt，输出旋转框列表。",
        "gui/result_window.py：校对总控、AI 线程、生成 HTML。",
        "gui/image_viewer.py：图层、命中、缩放平移。",
        "utils/html_generator.py：注入 accidentSettings 与 runtime。",
        "utils/liability_ai.py：上下文拼装与一次 API 调用。",
    ])
    d.card(s, 4.75, 1.4, 4.0, 5.2, "数据怎么走", [
        "航拍图路径 → YOLO detections。",
        "detections + 历史标注 → 校对场景。",
        "校对结果 → convert_to_3d_data（像素→米）。",
        "3D JSON 写入 templates/accident_scene.html。",
        "若启用 AI：同一套标注再生成 liability_context，后台线程请求。",
        "浏览器测量线 POST 回本机服务，另存 (line).html。",
    ])
    d.card(s, 8.95, 1.4, 3.8, 5.2, "路径解析", [
        "utils/app_paths.py：",
        "开发态用源码根目录；",
        "打包后用可执行文件目录。",
        "",
        "所有 config、downloads、history、templates、ExpertPic 都走 app_path()，避免相对路径在打包后失效。",
    ])

    s = d.new(title="仓库目录地图", subtitle="业务代码集中在 gui / core / utils / config")
    d.table(
        s, 0.55, 1.35, 12.2, 5.4,
        ["路径", "作用"],
        [
            ["main.py / start.bat / create.bat / build.bat", "入口、建环境、启动、打包"],
            ["gui/", "主窗、校对、运动编辑、系统设置各 Tab、尺寸/车道宽对话框"],
            ["core/", "YOLO 检测、8 种车型预设、5 种标记预设、车型规格读写"],
            ["utils/", "HTML 生成、像素→3D、AI、标注缓存、测量服务、动画/识别设置"],
            ["config/", "识别、AI 连接示例、提示词、专家、端口、车型默认/当前规格"],
            ["templates/accident_scene.html", "三维现场模板，注入 accidentSettings"],
            ["static/three/", "Three.js、OrbitControls、GLTFLoader（离线）"],
            ["downloads/", "YOLO .pt（仓库用 .gitkeep，权重大于 100MB 不上传）"],
            ["history/", "生成的 HTML 与测量副本"],
            ["ExpertPic/", "四位默认专家头像 + 自定义头像"],
            ["models/", "车型 GLB，导出时只内嵌本次用到的车型"],
        ],
        col_widths=[4.6, 7.6],
    )

    # 05 recognition
    d.section(3, "主流程与识别", "从选图到旋转框列表，再映射成可编辑的标准车型。")

    s = d.new(title="启动与识别线程", subtitle="识别在 QThread 中跑，避免卡住界面")
    d.bullets(s, [
        "启动：测量保存服务先 bind 端口；失败时主界面显示端口警告，查看历史/生成后的保存会受影响。",
        "上传图片：检查 downloads 中是否有可用 .pt；没有则提示先去「系统设置 → 识别设置」。",
        "ProcessThread 调用 YoloDetector(model_path).detect(image_path)，完成后打开 ResultWindow。",
        "若该图有标注历史（annotation fingerprint），校对页会恢复车辆、道路线、标记、路径、案情、本图车道宽。",
        "识别中主按钮禁用，进度条为不确定动画。",
        "识别出错弹窗后恢复按钮；权重缺失会抛 FileNotFoundError。",
    ])
    d.card(s, 0.55, 4.55, 12.2, 2.05, "YOLO 推理参数（core/yolo_detector.py）", [
        "conf=0.15，iou=0.4，imgsz=1024。有 OBB 时用 xywhr（中心、宽高、弧度角）；无 OBB 则退回水平框且 angle=0。",
        "每条检测含 class_id / class_name / confidence / x_center / y_center / width / height / angle。GUI 自己画框，draw_detections 仅调试用。",
    ])

    s = d.new(title="车辆类型如何从 YOLO 变成 3D 车型", subtitle="core/vehicle_presets.py：用户选定优先，否则按 class_id 推断")
    d.table(
        s, 0.55, 1.35, 12.2, 3.35,
        ["YOLO class_id", "推断车型", "备注"],
        [
            ["2 / 10", "小客车", "默认乘用车"],
            ["3", "商务车", ""],
            ["5", "大巴车", ""],
            ["7 / 9", "大型栏板货车", "卡货车类兜底也走栏板货"],
            ["名称含 large/truck/bus", "大型栏板货车", "class_id 未命中时的名字兜底"],
            ["名称含 small/car", "小客车", "名字兜底"],
        ],
        col_widths=[3.2, 3.2, 5.8],
    )
    d.card(s, 0.55, 4.9, 12.2, 1.7, "导出到 Three.js 的 8 种标准车型", [
        "小客车、商务车、牵引车及挂车、大型厢式货车、大型罐式货车、大型栏板货车、轻型栏板货车、大巴车。",
        "必须与 accident_scene.html 中 createVehicle 的 switch 一致。用户在校对里改的 preset_type 会覆盖自动推断。",
    ])

    s = d.new(title="像素坐标如何变成三维米制", subtitle="utils/data_converter.py：用车道宽度标定 ppm（像素/米）")
    d.bullets(s, [
        "默认车道宽 3.75 m、应急车道宽 3.0 m（config/vehicle_model_defaults.json，可在校对里按本图改）。",
        "根据标注的道路线间距与设定车道宽，计算 ppm，再把图像中心作为世界原点：x 向右、z 向前（图像 y）。",
        "道路线段两端各延长 30%（ROAD_MODEL_END_EXTENSION_RATIO），避免三维路面在画面边缘被切短。",
        "相邻道路线围成车道面：index 0 记为「应急车道」，其余为「快车道1、快车道2…」。",
        "车辆写入 type、位置、朝向、尺寸、车漆颜色（从航拍图采样）、运动路径（同步换算到米）与速度。",
        "标记物写入五种标准类型；碰撞点、散落物在运动播放期间隐藏，结束再出现。",
        "车身颜色解析失败则回退默认涂装。只把本次用到的车型 GLB 打进 HTML，控制文件体积。",
    ])

    # 06 editor
    d.section(4, "结果确认与编辑", "五图层校对是质量闸门：车辆和道路线不过关，不能生成三维现场。")

    s = d.new(title="校对页总览", subtitle="gui/result_window.py  +  gui/image_viewer.py")
    d.table(
        s, 0.55, 1.35, 12.2, 3.6,
        ["图层", "必填", "主要操作"],
        [
            ["*车辆", "是，至少 1 辆", "拖移、黄角点缩放、滚轮旋转、方向键转车头标识、Ctrl 多选、新增框"],
            ["*道路线", "是，至少 1 条", "拖整组平行线、滚轮旋转整组、新增、Delete/右键删最后一根"],
            ["标记物", "否", "单击放置圆标，右键改类型：椎桶/成年人/导向牌/散落物/碰撞点"],
            ["运动", "否", "右键车辆设路径/速度；路径点最多 10 个，滚轮切 G/F"],
            ["事故案情", "否", "浮动窗口输入简要案情，建议用车辆 ID 指代"],
        ],
        col_widths=[2.0, 2.2, 8.0],
    )
    d.card(s, 0.55, 5.15, 6.05, 1.45, "右侧主按钮", [
        "快捷键：红色呼吸闪烁，点开列出全部键鼠。",
        "仅缓存：写入标注历史，不生成 HTML。",
        "识别确认：车辆+道路线齐全才可点；可勾选本次启用 AI。",
    ])
    d.card(s, 6.8, 5.15, 5.95, 1.45, "显示全部图层", [
        "默认勾选。其它图层可见但不可点、不挡当前层命中。",
        "文字、箭头、路径、选框一律鼠标穿透。",
        "案情页会强制展示全部图层，方便用车辆 ID 写经过。",
    ])

    s = d.new(title="图层隔离与命中规则", subtitle="可见 ≠ 可操作；装饰物永远不能抢走鼠标")
    d.bullets(s, [
        "当前图层：对象可选、可拖、可右键。其它图层在「显示全部图层」时只渲染。",
        "未勾选时，非当前图层直接隐藏（运动页仍强制显示车辆框，便于选车设路径）。",
        "命中测试按当前模式过滤：道路线页点到车辆框不会误选车；车辆页点到车道线不会拖线。",
        "标签「车辆1(小客车)」字号约 6pt，贴旋转框左上角（YOLO 风格），不随图片缩放（ItemIgnoresTransformations）。",
        "道路线编号为绿色，与车道线颜色一致。",
        "非新增模式下点空白处平移：按下才抓握，松开必须恢复箭头，避免「拳头光标粘住」。",
        "设置路径时，车辆选框暂时不接收点击，允许把路径点打在车身上；退出路径模式后恢复选车。",
    ])

    s = d.new(title="车辆图层逻辑", subtitle="选框是现场事实，车型/尺寸/朝向是建模语义")
    d.card(s, 0.55, 1.4, 6.05, 5.2, "编辑能力", [
        "拖拽移动中心；拖黄色角点改宽高。",
        "选中后滚轮旋转整框（带旋转的航拍目标）。",
        "← / → 只旋转「车头方向标识」，不改框，用于告诉三维车头朝哪。",
        "Ctrl 点击多选，可批量转方向。",
        "右键：删除、改车型（8 种标准类型）、单独改该车尺寸。",
        "「新增」后拖拽画新框；再次点击退出新增。",
        "Delete 删除选中车辆。",
        "标签格式：车辆N(车型)，紧贴框左上。",
    ])
    d.card(s, 6.8, 1.4, 5.95, 5.2, "尺寸从哪来", [
        "全局默认：config/vehicle_model_settings.json（可由建模预览/设置改）。",
        "单车覆盖：校对里打开车辆尺寸对话框，只影响这一辆。",
        "导出时：单车尺寸 > 车型规格 > 代码兜底。",
        "牵引车及挂车按 tractorTrailer 特殊结构生成。",
        "车漆：从航拍图对应区域采样，让三维车身接近现场颜色。",
        "方向标识必须人工确认：YOLO 能给框的朝向，但不能可靠判断哪头是车头。",
    ])

    s = d.new(title="道路线图层逻辑", subtitle="道路线是三维地面姿态和车道划分的基准")
    d.bullets(s, [
        "一组平行线表示车道边界。拖动一根会保持平行间距，用于整体平移路网。",
        "选中后滚轮旋转整组，使灰地面/路模型与真实道路走向对齐，而不是平行于屏幕坐标轴。",
        "可设置本图车道宽、应急车道宽，写入标注缓存并参与 ppm 标定和 HTML 测量基准。",
        "相邻线之间自动分成应急车道 + 快车道，供 AI 上下文描述「车辆所在车道」。",
        "至少一条道路线才能点「识别确认」——没有路，三维地面无法定方向，测量锁定直角也没有参考轴。",
        "删除只能从最后一根往前删，避免中间抽掉导致车道面错乱。",
    ])

    s = d.new(title="标记物图层", subtitle="五种标准类型，导出到三维的几何不同")
    d.table(
        s, 0.55, 1.35, 12.2, 3.5,
        ["类型", "现场含义", "三维表现 / 播放行为"],
        [
            ["安全椎桶", "现场封控、分流", "锥桶模型，始终可见"],
            ["成年人", "行人、勘查人员等", "人物模型"],
            ["导向牌", "临时标志", "导向牌模型"],
            ["散落物", "玻璃、零件等", "地面散落；运动播放中隐藏，结束再现"],
            ["碰撞点", "接触位置", "黄色 3D 浮空扩散球；播放中隐藏，结束再现"],
        ],
        col_widths=[2.2, 3.5, 6.5],
    )
    d.card(s, 0.55, 5.1, 12.2, 1.5, "编辑要点", [
        "「新增」后左键单击连续放点；ID 自动为 M01、M02…。右键改类型。碰撞点与散落物是给讲解和 AI 用的语义标记，不是物理刚体。",
    ])

    s = d.new(title="运动图层：路径、速度、G/F", subtitle="gui/motion_path_editor.py  +  gui/motion_path_utils.py")
    d.card(s, 0.55, 1.4, 6.05, 5.2, "怎么画", [
        "仅运动页可选中车辆。右键「设置路径」进入绘制。",
        "左键空白（含车身）加点；最多 10 个用户点。",
        "右键空白结束：自动把终点连到该车中心（事故静止位置）。",
        "Delete 只能从后往前删当前呼吸点。",
        "滚轮切换当前点前进 G / 倒车 F，字母画在呼吸圆内且鼠标穿透。",
        "呼吸动画速度是普通闪烁的两倍，半径已减半，减少挡路。",
        "虚线在点与点之间流动，提示绘制顺序。",
        "车心旁蓝色标签显示预估时长，保留一位小数，例如 12.3秒。",
    ])
    d.card(s, 6.8, 1.4, 5.95, 5.2, "速度与颜色", [
        "右键「设置速度」，单位 km/h，默认 80。",
        "预估时长 = 路径米数 / (车速m/s × 系统执行倍率)。",
        "开场 1 秒静帧不计入这条预估，也不计入 HTML 进度条。",
        "超过 30 秒：路径变红并弹窗提醒，可提高车速或缩短路径。",
        "未超 30 秒：浅蓝色。",
        "G：车头指向下一路径点（前进）。",
        "F：车尾指向下一路径点（倒车）。车头转角连续，不在 G/F 切换时原地甩 180°。",
        "路径与速度随标注一起缓存。",
    ])

    s = d.new(title="事故案情与生成前校验", subtitle="案情是给 AI 和报告用的自然语言，不进入三维几何")
    d.bullets(s, [
        "事故案情页弹出可拖动浮动面板，建议用「车辆1」「车辆2」描述经过，避免写车牌造成隐私问题。",
        "「仅缓存」：保存车辆、道路、标记、路径、案情、本图车道宽，不建模。可反复改。",
        "「识别确认」禁用条件：没有车辆、或没有道路线。按钮变灰，悬停提示缺哪一项。下拉框里「车辆」「道路线」文字保持红色（不闪）。",
        "「本次启用AI分析」：仅当系统设置已勾选「启用AI分析责任」才可勾；否则禁用并提示先去系统设置。",
        "生成时：先落标注 → 可选后台 AI 线程（进度对话框）→ generate_html → 用系统浏览器打开。",
        "AI 失败不阻断三维生成：仍出 HTML，责任/报告按钮保持不可用或走已有缓存策略。",
        "「每次重新生成事故责任」开启时，忽略相同标注指纹的 AI 缓存，每次都打接口。",
    ])

    s = d.new(title="校对页快捷键一览", subtitle="与软件内「快捷键」弹窗一致")
    d.table(
        s, 0.55, 1.35, 12.2, 5.4,
        ["操作", "作用"],
        [
            ["Alt + 滚轮上 / 下", "放大 / 缩小图片；标签不跟着图片放大"],
            ["图片放大后左键拖空白 / 中键拖 / 空格+左键", "抓握平移视图"],
            ["Ctrl + 点击", "多选车辆"],
            ["选中车辆或道路线后滚轮", "旋转对象 / 整组车道线"],
            ["选中车辆后 ← / →", "旋转车头方向标识"],
            ["Delete 或右键", "删除当前对象；路径模式下删最后一个点"],
            ["新增后拖拽或单击", "按当前图层新增"],
            ["路径绘制中左键 / 右键 / 滚轮", "加点 / 结束并连车心 / 切换 G·F"],
        ],
        col_widths=[4.6, 7.6],
    )

    # 07 settings
    d.section(5, "系统设置全表", "六个 Tab 决定识别用哪套权重、要不要 AI、三维怎么播、测量怎么存。")

    s = d.new(title="系统设置地图", subtitle="gui/system_settings_dialog.py，模态对话框，底部可保存当前 Tab")
    d.table(
        s, 0.55, 1.35, 12.2, 5.4,
        ["Tab", "配置文件", "关键字段", "何时生效"],
        [
            ["识别设置", "recognition_settings.json", "modelFile（downloads 下的 .pt）", "下次上传识别"],
            ["AI连接", "ai_settings.json", "密钥、URL、模型、思考、超时、两个勾选", "识别确认且本次启用 AI"],
            ["AI提示词", "ai_prompt.json", "appendRequirement（交警视角补句）", "拼进当次请求"],
            ["AI专家设置", "ai_experts.json + ExpertPic", "最多 12 人，角色描述 ≤500 字", "当次 AI 请求的 experts 列表"],
            ["端口设置", "measurement_port.json", "silentRandom、port（默认 8765）", "保存后立即重启监听"],
            ["动画设置", "animation_settings.json", "playbackSpeedFactor 0.02–0.25，默认 0.1", "写入新生成 HTML 的 runtime"],
        ],
        col_widths=[2.2, 3.3, 4.0, 2.7],
    )

    s = d.new(title="识别设置", subtitle="只做一件事：选定 YOLO 权重文件名")
    d.bullets(s, [
        "下拉框扫描 downloads/ 中全部 .pt。一个都没有时禁用保存，提示未找到模型。",
        "默认优先选文件名不含 obb 的权重；用户仍可手动选 yolo11x-obb.pt 以获得旋转框。",
        "推荐组合：yolo11x.pt（水平框）与 yolo11x-obb.pt（旋转框）。仓库不包含这两个文件。",
        "设置保存在 config/recognition_settings.json 的 modelFile。",
        "已打开的校对页不会中途换模型；必须重新上传图片才会按新权重识别。",
    ])

    s = d.new(title="AI 连接", subtitle="未勾选「启用AI分析责任」或关键字段为空，全程不发 HTTP")
    d.table(
        s, 0.55, 1.35, 12.2, 5.4,
        ["字段", "默认 / 范围", "说明"],
        [
            ["API密钥", "空，密码框", "仅存本机 ai_settings.json；打包与 Git 都会去掉"],
            ["请求地址", "https://api.deepseek.com/chat/completions", "必须 http/https 开头"],
            ["模型名称", "deepseek-v4-pro", "随服务商填写"],
            ["推理强度", "high（低/中/高/默认）", "reasoningEffort"],
            ["思考模式", "enabled", "default / enabled / disabled"],
            ["Temperature", "0.1（0–2）", "责任分析需要低随机性"],
            ["超时时间", "120 秒（10–600）", "专家发言一次性返回，超时宜留足"],
            ["启用AI分析责任", "false", "总闸。关掉则校对页「本次启用AI」也禁用"],
            ["每次重新生成事故责任", "true", "true=忽略同指纹缓存，每次都请求"],
        ],
        col_widths=[3.4, 4.0, 4.8],
    )

    s = d.new(title="AI 提示词与专家设置", subtitle="提示词补交警视角；专家决定报告里谁按什么口吻说话")
    d.card(s, 0.55, 1.4, 6.05, 5.2, "AI提示词", [
        "文件：config/ai_prompt.json。",
        "字段 appendRequirement 会追加到当次 prompt.requirements。",
        "默认要求：站在交警视角谈防控，",
        "谈如何修路、如何打击违法，",
        "而不是教驾驶人「以后怎么开」。",
        "可改，但应保持可执行、可部署的机关口吻。",
    ])
    d.card(s, 6.8, 1.4, 5.95, 5.2, "AI专家设置规则", [
        "最多 12 位。满员后「+」禁用。",
        "角色定位最多 500 字。",
        "头像可选；若上传须 1:1 正方形且 ≤500KB。",
        "未上传则用问号默认图 default_unknown.png。",
        "四位默认专家：不可删除、不可停用，可改角色描述。",
        "自定义专家：可改、可删、可停用；发言顺序=创建顺序，排在四位默认之后。",
        "仅 enabled 的专家进入当次请求。",
        "头像在 HTML 里 base64 内嵌，离线也能显示。",
    ])

    s = d.new(title="四位默认专家（固定发言顺序）", subtitle="案件侦办 → 秩序管理 → 护路联防 → 指挥调度")
    d.card(s, 0.55, 1.4, 6.05, 2.45, "案件侦办大队长  ·  案侦.png", [
        "还原经过：车道、车速、方向、碰撞点、护栏受损。",
        "判断主要 / 次要 / 同等 / 无责任，并给出下一步调查方向。",
    ])
    d.card(s, 6.8, 1.4, 5.95, 2.45, "秩序管理大队长  ·  秩序.png", [
        "点名各车违法行为及对事故的影响权重。",
        "提出执法整治专项与警示宣传角度。",
    ])
    d.card(s, 0.55, 4.05, 6.05, 2.45, "护路联防大队长  ·  护路.png", [
        "看道路因素：应急车道过窄、摩擦系数、标志标牌、抓拍、提示手段、减速设施。",
        "建议路面应增设什么。",
    ])
    d.card(s, 6.8, 4.05, 5.95, 2.45, "指挥调度大队长  ·  指挥.png", [
        "盯监控/GPS 能发现的动态风险：违停、倒车、占应急、行人上高速。",
        "给出指挥中心发现与制止的手段清单。",
    ])

    s = d.new(title="端口设置与动画设置", subtitle="端口服务给 HTML 存文件；动画系数给运动播放定快慢")
    d.card(s, 0.55, 1.4, 6.05, 5.2, "端口设置", [
        "测量保存是进程内 HTTP 服务，HTML 通过 localhost POST。",
        "默认端口 8765。",
        "勾选「静默随机端口」：自动选可用端口（含历史别名与随机回退）。",
        "取消勾选：只用指定端口；占用则保存后服务起不来。",
        "「测试」探测端口是否可用（本进程占用视为正常）。",
        "「恢复默认」回到静默随机 + 8765。",
        "保存后立即 restart_measurement_save_server()。",
        "服务未启动时，HTML 会提示先开 Python 程序再保存测量线。",
        "只允许把 history 目录内的 HTML 另存，防止写到任意路径。",
    ])
    d.card(s, 6.8, 1.4, 5.95, 5.2, "动画设置", [
        "字段 playbackSpeedFactor。",
        "范围 0.02–0.25，步进 0.01，默认 0.1。",
        "含义：三维车速 = 校对里填的 km/h × 该系数。",
        "0.1 即按真实车速的十分之一播放，便于看清会车与转向。",
        "该系数写入新 HTML 的 accidentSceneRuntime。",
        "HTML 里另有 1X / 2X / 3X，叠乘在该系数之上。",
        "预估时长、进度条长度都用「行驶时间」，不含开场 1 秒静帧。",
    ])

    s = d.new(title="车型规格与建模预览", subtitle="三维尺寸来自 JSON，不在 C++/着色器里写死")
    d.bullets(s, [
        "默认规格：config/vehicle_model_defaults.json；当前覆盖：vehicle_model_settings.json。",
        "每车含 kind（passenger / tractorTrailer / boxTruck 等）、宽长高、轮半径、轮位纵向偏移。",
        "全局车道宽 3.75 m、应急 3.0 m、描边宽 2.5，可被单图校对覆盖。",
        "主窗口「建模预览」打开 templates/vehicle_model_preview.html，底部三角切换 8 种车型，逻辑与导出场景同源。",
        "改规格后，新生成的 HTML 用新尺寸；已生成的旧 HTML 不会自动改。",
    ])

    # 08 HTML
    d.section(6, "三维报告与播放", "浏览器里的现场：能转、能量、能播、能讲责任。")

    s = d.new(title="HTML 生成链路", subtitle="utils/html_generator.py 读模板、注入 JSON、只打包用到的 GLB")
    d.process(
        s,
        [
            ("读航拍图", "取宽高，供坐标换算与车漆采样"),
            ("转 3D JSON", "vehicles / roads / markers / motion / laneAreas"),
            ("灌模板", "替换 settings.js 为内联 accidentSettings"),
            ("写 runtime", "源路径、保存 URL、播放倍率、专家头像"),
        ],
        y=1.45,
        box_h=1.3,
    )
    d.bullets(s, [
        "输出：history/{YYYYMMDD_HHMMSS}.html，用系统浏览器打开。",
        "runtime 含 sourceHtmlPath、saveMeasurementsUrl、playbackSpeedFactor；测量线以后另存时靠它找回源文件。",
        "Three.js 与控制器、加载器均已放进 static/three，生成页可离线打开（保存测量除外，保存仍要本机服务）。",
        "灰地面按道路线走向旋转，与路平行，而不是平行于世界坐标轴。",
    ], top=3.05, height=3.5)

    s = d.new(title="三维现场交互总表", subtitle="templates/accident_scene.html 右侧按钮与三个勾选框")
    d.table(
        s, 0.55, 1.35, 12.2, 5.4,
        ["控件", "默认", "互斥 / 禁用规则"],
        [
            ["开场建模动画", "自动播", "期间业务按钮全禁用；任意左键立即结束并解锁"],
            ["测量距离", "关", "测量中禁用运动/责任/报告/保存；三个勾选框仍可用；再点退出"],
            ["保存结果", "需服务", "弹出文件名；最终 history/{名}_{时间}(line).html；非法字符拒绝；重名确认"],
            ["播放运动", "关", "无路径则禁用；责任模式下禁用；播放中隐藏测量线与责任 UI"],
            ["责任显示", "关", "无 AI 结构化结果则保持灰；开启时隐藏车辆编号"],
            ["事故报告", "关", "只渲染「头像+姓名+发言」；旧版总述/清单/建议不再展示"],
            ["磁吸模式", "关，不保存", "吸附轮胎、标记中心、道路线"],
            ["锁定直角", "关，不保存", "相对车道线 0°/90°；可与磁吸叠加到延长线最近点"],
            ["车辆编号", "关，不保存", "蓝色头顶标牌，样式同责任标牌；测量/运动/责任时隐藏"],
        ],
        col_widths=[2.2, 2.0, 8.0],
    )

    s = d.new(title="测量距离", subtitle="建模设置里的车道宽、应急车道宽是测距的比例尺")
    d.bullets(s, [
        "进入测量：点路面第一点为起点，移动时虚线预览距离与相对车道线夹角，再点第二点落线。",
        "磁吸：终点吸到最近轮胎、标记中心或道路线。",
        "锁定直角：终点被约束到与车道线平行或垂直的方向。",
        "两者同时开：先正交，再在该延长线上吸到最近特征点。",
        "保存：校验文件名不含 \\ / : * ? \" < > | ；时间戳格式 _{YYYYMMDD_HHMMSS}(line).html。",
        "服务把测量线写进副本 HTML 的 runtime.savedMeasurements，下次打开仍在。",
        "未启动桌面程序时不能保存，页面提示先启动 Python 程序。",
    ])

    s = d.new(title="运动播放算法", subtitle="所有已设路径的车辆同时从起点出发，车头连续转向")
    d.card(s, 0.55, 1.4, 6.05, 5.2, "时间轴", [
        "进入播放：车辆跳到各自路径起点。",
        "先静持 1 秒（不是 2 秒），进度条仍为 0。",
        "然后按各自 km/h × 系统系数 × 界面 1X/2X/3X 沿路径前进。",
        "进度条只覆盖「行驶段」，可拖动；一拖就暂停。",
        "播放/暂停按钮控制时钟。",
        "散落物与碰撞点在行驶期间隐藏，结束或退出后出现。",
        "退出播放：车辆回到事故静止位（路径终点/车心），恢复其它按钮。",
        "无路径车辆保持事故位置不动。",
    ])
    d.card(s, 6.8, 1.4, 5.95, 5.2, "姿态与平滑", [
        "路径折线在 HTML 里做圆弧倒角（fillet），小转角也保证看得见的圆角，避免折线「掰一下」。",
        "G：车头对准切线方向；F：车头与切线相反（倒车）。",
        "约束的是「每帧车头角变化」，不是路径几何角，也不是 G/F 混合。",
        "最后几米把 yaw blend 到校对时的 restYaw，最后一帧车头等于标注方向，无猛地一扭。",
        "倒车时车头仍连续，不会为了切 F 而原地掉头。",
    ])

    s = d.new(title="责任显示与事故报告", subtitle="同一份 AI JSON，两套界面：车上标牌 vs 专家发言列表")
    d.bullets(s, [
        "一次 DeepSeek 请求必须同时返回 vehicleResponsibilities 与 expertSpeeches，禁止按专家拆多次。",
        "责任显示：每车头顶标牌（主要/次要/同等/无责任等），悬停看该车说明。开启时车辆编号隐藏。",
        "事故报告：按 experts 数组顺序逐人展示头像、姓名、第一人称发言。默认四人固定在前，自定义按创建序接在后面。",
        "报告面板不再展示「总体综述 / 责任清单 / 防控建议」旧三块，避免和专家口吻重复。",
        "代码仍会做责任一致性修补（_enforce_responsibility_consistency），减少自相矛盾的划责。",
        "未启用 AI 或请求失败：三维现场仍可测、可播；责任与报告按钮保持灰色。",
    ])

    s = d.new(title="AI 分析数据流（一次请求）", subtitle="utils/liability_ai.py：先把现场编成可观察摘要，再交给模型")
    d.card(s, 0.55, 1.4, 12.2, 2.3, "上下文里有什么（校对页全部图层都会进）", [
        "车辆摘要：编号、车型、所在车道、朝向、与其它车的碰撞/邻近关系、运动路径与速度、G/F。",
        "道路：车道数量与应急/快车道划分、线型。标记：椎桶、行人、导向牌、散落物、碰撞点位置。",
        "案情原文。自定义提示词 appendRequirement。experts[]：id、name、rolePrompt（仅启用项）。",
    ])
    d.card(s, 0.55, 3.9, 12.2, 2.6, "返回约定", [
        "vehicleResponsibilities：按车辆排序的结构化划责，供「责任显示」。",
        "expertSpeeches：与 experts 顺序一一对应的第一人称发言，供「事故报告」。",
        "温度默认 0.1；思考模式默认开启。后台 QThread 跑，界面弹进度框，不冻校对窗。",
        "相同标注指纹可缓存；「每次重新生成」开启则跳过缓存。",
    ])

    # 09 data security
    d.section(7, "数据、打包与安全", "配置可分发，密钥不可离开这台机器。")

    s = d.new(title="配置与缓存文件", subtitle="业务 JSON 要进安装包；真正的密钥文件不进 Git")
    d.table(
        s, 0.55, 1.35, 12.2, 5.4,
        ["文件", "进 Git?", "进安装包?", "说明"],
        [
            ["config/ai_settings.json", "否（gitignore）", "是，但密钥被清空", "pack_sanitize_ai_settings.py 擦除 deepseekApiKey"],
            ["config/ai_settings.example.json", "是", "可选", "空密钥模板，给开源用户复制"],
            ["ai_prompt.json / ai_experts.json", "是", "是", "提示词与专家角色"],
            ["knowledge/playbook.json", "是", "是", "全库唯一口径手册"],
            ["knowledge/index.json 与 packs/*.json", "是", "是", "精炼案件结果列表"],
            ["recognition / animation / measurement_port / vehicle_model_*.json", "是", "是", "识别、动画、端口、车型"],
            ["downloads/*.pt", "否", "是", "GitHub 100MB 限制；打包给现场离线用"],
            ["history/*.html", "否", "否", "案件产物，不进仓库"],
            ["ExpertPic 默认 PNG", "是", "是", "四位大队长头像"],
        ],
        col_widths=[4.4, 2.0, 2.2, 3.6],
    )

    s = d.new(title="标注缓存", subtitle="utils/annotation_history.py：同一张图再次打开，接着改而不是从头画")
    d.bullets(s, [
        "用图片路径与内容指纹关联缓存，避免「换了一张同名图却套用旧框」。",
        "缓存内容：车辆框与车型尺寸、道路线、标记、运动路径与速度、事故案情、本图车道宽。",
        "「仅缓存」与「识别确认」都会写；识别确认额外生成 HTML。",
        "AI 结果可按标注指纹复用，除非设置了每次重新生成。",
    ])

    s = d.new(title="打包逻辑（build.bat）", subtitle="产物在 release/TrafficAccident3D/，不使用 dist 作为最终交付")
    d.bullets(s, [
        "使用项目内 Conda 环境 env/，Python 3.10。依赖走清华 PyPI 镜像。",
        "PyInstaller 打主程序；再拷贝 config、templates、static、models、downloads、ExpertPic、knowledge 等到 release。",
        "拷贝 AI 设置前先跑 pack_sanitize_ai_settings.py，保证安装包里 API 密钥为空。",
        "knowledge 会带上口径手册和精炼案件结果；不拷 Excel 源表和 .migrated 备份。",
        "pack_verify_release.py 检查：必要配置在、YOLO 权重在、知识库文件在、包装后的 ai_settings.json 密钥为空、整个 release 目录搜不到开发机上的原密钥。",
        "BUILD_NOPAUSE=1 可无人值守。失败统一 goto :fail。",
        "dist/ 是 PyInstaller 中间目录，最终以 release 为准。",
    ])

    s = d.new(title="开源与密钥隔离", subtitle="公开仓库 https://github.com/crawar/TrafficAccident3D")
    d.card(s, 0.55, 1.4, 6.05, 5.2, ".gitignore 关键项", [
        "env/、venv/、build/、dist/、release/",
        "config/ai_settings.json",
        "history/*（保留 .gitkeep）",
        "knowledge 下的 Excel、tmp、.migrated",
        "downloads/*.pt / .onnx / .engine",
        ".cursor/ 等编辑器目录",
        "",
        "publish_github.bat 在 push 前会 git check-ignore 密钥文件，并拒绝已暂存的 ai_settings.json。",
    ])
    d.card(s, 6.8, 1.4, 5.95, 5.2, "使用者怎么配 AI", [
        "复制 ai_settings.example.json 为 ai_settings.json，",
        "或在软件「系统设置 → AI连接」里填写。",
        "不配密钥：识别、校对、三维、测量、运动全部可用。",
        "YOLO 权重需自行放到 downloads/：",
        "yolo11x.pt 、 yolo11x-obb.pt。",
        "README 只写功能与步骤，不写任何真实密钥。",
    ])

    s = d.new(title="能力边界与常见误用", subtitle="先知道边界，再向领导和兄弟单位介绍")
    d.table(
        s, 0.55, 1.35, 12.2, 5.4,
        ["场景", "系统能做什么", "不能当成什么"],
        [
            ["责任认定", "提供可视化现场 + 专家口吻分析草稿", "不能代替法定事故认定书"],
            ["测距", "按标定车道宽在三维里量相对距离", "不是全站仪级测绘成果"],
            ["运动回放", "按民警标注的路径与速度示意经过", "不是从视频里自动追轨迹"],
            ["YOLO 识别", "给出初始旋转框，大幅减少手画", "必须人工确认车头与漏检"],
            ["AI 发言", "一次生成四岗+自定义专家讲稿", "未开启或失败时报告按钮为灰"],
            ["离线演示", "HTML 可拷走看场景和播放", "保存测量线仍需本机程序在跑"],
        ],
        col_widths=[2.2, 5.0, 5.0],
    )

    s = d.new(title="标准操作清单（给现场民警）", subtitle="五步走完一件案子")
    d.process(
        s,
        [
            ("准备", "create.bat 建 env；downloads 放两个 .pt；可选填密钥"),
            ("识别", "上传航拍；等 YOLO；进入校对"),
            ("校对", "车与路必填；补标记、路径、案情"),
            ("生成", "识别确认；浏览器打开 HTML"),
            ("讲解", "测量 / 编号 / 播放 / 责任 / 报告 / 另存"),
        ],
        y=1.5,
        box_h=1.4,
    )
    d.card(s, 0.55, 3.3, 12.2, 3.2, "检查口诀", [
        "车头方向看箭头，不看框的长边。道路线要转到和路平行。应急车道是最外侧那一条带。",
        "路径点从来车方向画到事故停车位置；倒车路段用 F。超过 30 秒的红线先改速度。",
        "要讲责任必须先在系统设置打开 AI，并在识别确认时勾选本次启用。报告里发言顺序不能在 HTML 里改，要在专家设置里改。",
        "把 HTML 拷给领导可以看；若还要在那台机器上保存测量线，需要同时运行本程序且端口通。",
    ])

    # closing
    s = d.new(title="模块对照速查", subtitle="改需求时先对上文件，避免在 HTML 和 Python 各改一半")
    d.table(
        s, 0.55, 1.35, 12.2, 5.4,
        ["要改的行为", "优先看这些文件"],
        [
            ["主界面按钮与识别线程", "gui/main_window.py，core/yolo_detector.py"],
            ["校对图层、必填、快捷键、AI 勾选", "gui/result_window.py，gui/image_viewer.py"],
            ["路径点、G/F、时长颜色", "gui/motion_path_editor.py，motion_path_utils.py"],
            ["系统六个 Tab", "gui/system_settings_dialog.py 及对应 utils/*_settings.py"],
            ["专家增删改头像", "gui/ai_expert_settings_dialog.py，utils/ai_experts.py"],
            ["像素→米、车道面、车漆", "utils/data_converter.py，core/vehicle_presets.py"],
            ["三维按钮、测量、播放、标牌", "templates/accident_scene.html，utils/html_generator.py"],
            ["AI 上下文与返回修补", "utils/liability_ai.py，config/ai_prompt.json"],
            ["测量保存与文件名", "utils/measurement_save_server.py，html_generator.save_measurement_html"],
            ["打包与密钥擦除", "build.bat，pack_sanitize_ai_settings.py，pack_verify_release.py"],
        ],
        col_widths=[4.4, 7.8],
    )

    s = d.new(accent=False, footer=False)
    band = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.22), SLIDE_H)
    _fill(band, TEAL)
    _no_line(band)
    _textbox(s, Inches(0.8), Inches(2.2), Inches(11.5), Inches(0.4), "ACCIDENT MIND  ·  AI", 14, TEAL_LIGHT, True)
    _textbox(s, Inches(0.8), Inches(2.75), Inches(11.8), Inches(1.1), "事故智脑：人工智能与交通事故的结合。", 28, WHITE, True)
    _textbox(
        s, Inches(0.8), Inches(4.1), Inches(11.5), Inches(1.2),
        "本 PPT 由 PPT/generate_ppt.py 根据当前代码生成。\n修改逻辑后重新运行脚本即可更新，无需手改每一页。",
        16, MUTED,
    )
    d.chips(s, ["无密钥仍可建模", "车辆+道路线必填", "AI 一次请求全角色", "密钥不上仓库"], 0.8, 5.5, 2.7, 0.45)

    d.finish(OUT_PATH)
    return OUT_PATH, len(d.slides)


if __name__ == "__main__":
    path, n = build()
    print(f"Wrote {n} slides -> {path}")
