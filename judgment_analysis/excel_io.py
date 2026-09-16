# -*- coding: utf-8 -*-
"""Load .xlsx / .xls, expand merged cells, return raw rows for header mapping."""

import os
from datetime import date, datetime, time

SUPPORTED_EXCEL = {".xlsx", ".xls"}
SCAN_COLS = 3
SUSPICIOUS_WORDS = ("统计", "合计", "总计")
HEADER_CANDIDATE_LIMIT = 80


class ExcelLoadError(Exception):
    pass


def _is_empty(value):
    if value is None:
        return True
    if isinstance(value, str) and not value.strip():
        return True
    return False


def _normalize_value(value):
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value
    if isinstance(value, date):
        return value
    if isinstance(value, time):
        return value
    if isinstance(value, bool):
        return value
    if isinstance(value, float):
        if value != value:
            return ""
        if value == int(value) and abs(value) < 1e15:
            return int(value)
        return value
    if isinstance(value, int):
        return value
    text = str(value).strip()
    return text


def _notify(progress_cb, percent, message):
    if progress_cb:
        progress_cb(max(0, min(100, int(percent))), message)


def _trim_grid(grid):
    while grid and all(_is_empty(cell) for cell in grid[-1]):
        grid.pop()
    if not grid:
        return []
    last_col = 0
    for row in grid:
        for index, cell in enumerate(row):
            if not _is_empty(cell):
                last_col = max(last_col, index + 1)
    if last_col <= 0:
        return []
    return [list(row[:last_col]) for row in grid]


def uniquify_headers(raw_headers):
    seen = {}
    headers = []
    for index, item in enumerate(raw_headers):
        name = str(item).strip() if not _is_empty(item) else ""
        if not name:
            name = "列%d" % (index + 1)
        count = seen.get(name, 0)
        if count:
            headers.append("%s (%d)" % (name, count + 1))
        else:
            headers.append(name)
        seen[name] = count + 1
    return headers


def _drop_empty_rows(rows):
    return [row for row in rows if any(not _is_empty(cell) for cell in row)]


def _pad_row(row, width):
    padded = list(row[:width])
    if len(padded) < width:
        padded.extend([""] * (width - len(padded)))
    return padded


def apply_header_row(raw_rows, header_row_index):
    if not raw_rows:
        raise ExcelLoadError("工作表为空，没有可读的表头或数据。")
    if header_row_index is None or header_row_index < 0 or header_row_index >= len(raw_rows):
        raise ExcelLoadError("请先选择表头所在行。")
    width = max(len(row) for row in raw_rows)
    headers = uniquify_headers(_pad_row(raw_rows[header_row_index], width))
    rows = []
    for raw in raw_rows[header_row_index + 1 :]:
        rows.append(_pad_row(raw, width))
    rows = _drop_empty_rows(rows)
    return headers, rows


def resolve_header_row(raw_rows, saved_index, saved_labels):
    """Reuse last header mapping when the same row or the same labels are found."""
    if not raw_rows:
        return None
    width = max(len(row) for row in raw_rows)
    labels = [str(item) for item in (saved_labels or [])]

    def labels_at(index):
        return uniquify_headers(_pad_row(raw_rows[index], width))

    if labels:
        if saved_index is not None and 0 <= int(saved_index) < len(raw_rows):
            if labels_at(int(saved_index)) == labels:
                return int(saved_index)
        limit = min(len(raw_rows), HEADER_CANDIDATE_LIMIT)
        for index in range(limit):
            if labels_at(index) == labels:
                return index
    if saved_index is not None and 0 <= int(saved_index) < len(raw_rows):
        return int(saved_index)
    return None


def scan_suspicious_rows(grid):
    """Flag rows whose first SCAN_COLS cells contain 统计 / 合计 / 总计."""
    flagged = []
    for row_index, row in enumerate(grid or []):
        limit = min(SCAN_COLS, len(row))
        for col in range(limit):
            text = str(row[col] if row[col] is not None else "")
            if any(word in text for word in SUSPICIOUS_WORDS):
                flagged.append(row_index + 1)
                break
    return flagged


def header_row_preview(row, max_cells=6):
    parts = []
    for cell in list(row or [])[:max_cells]:
        text = "" if _is_empty(cell) else str(cell).replace("\n", " ").strip()
        if len(text) > 16:
            text = text[:16] + "…"
        parts.append(text or "（空）")
    return " ／ ".join(parts) if parts else "（空行）"


def _fill_merges(grid, ranges):
    for min_row, max_row, min_col, max_col in ranges:
        if min_row < 0 or min_row >= len(grid):
            continue
        row = grid[min_row]
        if min_col < 0 or min_col >= len(row):
            value = ""
        else:
            value = row[min_col]
        for row_index in range(min_row, min(max_row + 1, len(grid))):
            current = grid[row_index]
            needed = max_col + 1 - len(current)
            if needed > 0:
                current.extend([""] * needed)
            for col in range(min_col, max_col + 1):
                if 0 <= col < len(current):
                    current[col] = value


def _pack_result(names, active_name, target, grid, merges, progress_cb=None):
    _notify(progress_cb, 88, "正在检查可能的统计行")
    suspicious = scan_suspicious_rows(grid)
    _notify(progress_cb, 100, "载入完成")
    return {
        "sheet_names": names,
        "active_name": active_name,
        "sheet_name": target,
        "raw_rows": grid,
        "merges": list(merges or []),
        "suspicious_rows": suspicious,
    }


def _load_xlsx(path, sheet_name=None, progress_cb=None):
    from openpyxl import load_workbook

    _notify(progress_cb, 6, "正在打开工作簿")
    workbook = load_workbook(path, data_only=True, read_only=False)
    try:
        names = list(workbook.sheetnames)
        if not names:
            raise ExcelLoadError("工作簿中没有工作表。")
        active_name = workbook.active.title if workbook.active is not None else names[0]
        target = sheet_name or active_name
        if target not in workbook.sheetnames:
            raise ExcelLoadError("找不到工作表：%s" % target)
        sheet = workbook[target]
        max_row = sheet.max_row or 0
        max_col = sheet.max_column or 0
        if max_row < 1 or max_col < 1:
            raise ExcelLoadError("工作表为空，没有可读的表头或数据。")
        _notify(progress_cb, 12, "正在读取单元格")
        grid = []
        step = 50 if max_row > 400 else 10
        for index, row in enumerate(
            sheet.iter_rows(
                min_row=1, max_row=max_row, min_col=1, max_col=max_col, values_only=True
            ),
            start=1,
        ):
            grid.append([_normalize_value(cell) for cell in row])
            if index == 1 or index == max_row or index % step == 0:
                percent = 12 + int(58 * index / max(max_row, 1))
                _notify(progress_cb, percent, "正在读取表格（%d/%d）" % (index, max_row))
        merges = []
        for merged in sheet.merged_cells.ranges:
            min_col, min_row, max_col, max_row_bound = merged.bounds
            merges.append((min_row - 1, max_row_bound - 1, min_col - 1, max_col - 1))
        _notify(progress_cb, 76, "正在展开合并单元格")
        _fill_merges(grid, merges)
        grid = _trim_grid(grid)
        if not grid:
            raise ExcelLoadError("工作表为空，没有可读的表头或数据。")
        return _pack_result(names, active_name, target, grid, merges, progress_cb)
    finally:
        workbook.close()


def _xls_cell_value(sheet, row, col, datemode):
    import xlrd

    cell = sheet.cell(row, col)
    if cell.ctype == xlrd.XL_CELL_EMPTY or cell.ctype == xlrd.XL_CELL_BLANK:
        return ""
    if cell.ctype == xlrd.XL_CELL_DATE:
        try:
            return xlrd.xldate_as_datetime(cell.value, datemode)
        except Exception:
            return _normalize_value(cell.value)
    if cell.ctype == xlrd.XL_CELL_BOOLEAN:
        return bool(cell.value)
    if cell.ctype == xlrd.XL_CELL_ERROR:
        return ""
    return _normalize_value(cell.value)


def _load_xls(path, sheet_name=None, progress_cb=None):
    import xlrd

    _notify(progress_cb, 6, "正在打开工作簿")
    try:
        book = xlrd.open_workbook(path, formatting_info=True)
    except Exception:
        book = xlrd.open_workbook(path)
    names = list(book.sheet_names())
    if not names:
        raise ExcelLoadError("工作簿中没有工作表。")
    active_name = names[0]
    target = sheet_name or active_name
    if target not in names:
        raise ExcelLoadError("找不到工作表：%s" % target)
    sheet = book.sheet_by_name(target)
    nrows = sheet.nrows
    ncols = sheet.ncols
    if nrows < 1 or ncols < 1:
        raise ExcelLoadError("工作表为空，没有可读的表头或数据。")
    _notify(progress_cb, 12, "正在读取单元格")
    grid = []
    step = 50 if nrows > 400 else 10
    for row in range(nrows):
        grid.append(
            [_xls_cell_value(sheet, row, col, book.datemode) for col in range(ncols)]
        )
        index = row + 1
        if index == 1 or index == nrows or index % step == 0:
            percent = 12 + int(58 * index / max(nrows, 1))
            _notify(progress_cb, percent, "正在读取表格（%d/%d）" % (index, nrows))
    merges = []
    for rlo, rhi, clo, chi in getattr(sheet, "merged_cells", []) or []:
        merges.append((rlo, rhi - 1, clo, chi - 1))
    _notify(progress_cb, 76, "正在展开合并单元格")
    _fill_merges(grid, merges)
    grid = _trim_grid(grid)
    if not grid:
        raise ExcelLoadError("工作表为空，没有可读的表头或数据。")
    return _pack_result(names, active_name, target, grid, merges, progress_cb)


def load_sheet(path, sheet_name=None, progress_cb=None):
    path = os.path.abspath(path)
    if not os.path.isfile(path):
        raise ExcelLoadError("所选文件不存在。")
    ext = os.path.splitext(path)[1].lower()
    if ext not in SUPPORTED_EXCEL:
        raise ExcelLoadError("仅支持 Excel（.xlsx、.xls）。")
    _notify(progress_cb, 2, "正在准备载入")
    if ext == ".xlsx":
        data = _load_xlsx(path, sheet_name, progress_cb=progress_cb)
    else:
        data = _load_xls(path, sheet_name, progress_cb=progress_cb)
    data["path"] = path
    return data
