# -*- coding: utf-8 -*-
"""Build the A4 accident analysis Word report."""

import os
from datetime import datetime

from judgment_analysis.charts import bar_chart, hour_line_chart, pie_chart
from judgment_analysis.stats import build_report_data

NO_DATA = "未找到相关数据。"
CANNOT_ANALYZE = "未映射事故原因，无法分析。"


class ReportCancelled(Exception):
    pass


class ReportError(Exception):
    pass


def _raise_if_cancelled(should_cancel):
    if should_cancel and should_cancel():
        raise ReportCancelled("已取消研判分析。")


def output_path_for(src_path, when=None):
    stamp = (when or datetime.now()).strftime("%Y%m%d_%H%M")
    src_path = os.path.abspath(src_path or "")
    folder = os.path.dirname(src_path)
    if not folder or not os.path.isdir(folder):
        raise ReportError("无法确定源文件所在目录，未写入报告。")
    name = "事故分析报告_%s.docx" % stamp
    dest = os.path.join(folder, name)
    if not os.path.isfile(dest):
        return dest
    index = 2
    while True:
        dest = os.path.join(folder, "事故分析报告_%s(%d).docx" % (stamp, index))
        if not os.path.isfile(dest):
            return dest
        index += 1


def _set_run_font(run, font_name, size_pt, bold=False):
    from docx.oxml.ns import qn
    from docx.shared import Pt

    run.bold = bold
    run.font.size = Pt(size_pt)
    run.font.name = font_name
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn("w:eastAsia"), font_name)
    r_fonts.set(qn("w:ascii"), font_name)
    r_fonts.set(qn("w:hAnsi"), font_name)


def _add_paragraph(document, text, font_name="仿宋", size_pt=16, bold=False, align=None):
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    paragraph = document.add_paragraph()
    if align is not None:
        paragraph.alignment = align
    else:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_after = None
    run = paragraph.add_run(text)
    _set_run_font(run, font_name, size_pt, bold=bold)
    return paragraph


def _add_title(document, text):
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    return _add_paragraph(
        document, text, font_name="黑体", size_pt=22, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER
    )


def _add_h1(document, text):
    return _add_paragraph(document, text, font_name="黑体", size_pt=16, bold=True)


def _add_h2(document, text):
    return _add_paragraph(document, text, font_name="楷体", size_pt=16, bold=False)


def _add_body(document, text):
    return _add_paragraph(document, text, font_name="仿宋", size_pt=16, bold=False)


def _set_cell_text(cell, text, header=False):
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(str(text))
    if header:
        _set_run_font(run, "黑体", 12, bold=True)
    else:
        _set_run_font(run, "仿宋", 12, bold=False)


def _set_table_grid(table):
    table.style = "Table Grid"
    table.autofit = True
    try:
        from docx.enum.table import WD_TABLE_ALIGNMENT

        table.alignment = WD_TABLE_ALIGNMENT.CENTER
    except Exception:
        pass


def _add_table(document, headers, rows):
    table = document.add_table(rows=1 + len(rows), cols=len(headers))
    _set_table_grid(table)
    for index, header in enumerate(headers):
        _set_cell_text(table.rows[0].cells[index], header, header=True)
    for row_index, row in enumerate(rows):
        for col_index, value in enumerate(row):
            _set_cell_text(table.rows[row_index + 1].cells[col_index], value, header=False)
    return table


def _add_picture(document, image_path):
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm

    if not image_path or not os.path.isfile(image_path):
        return
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.add_picture(image_path, width=Cm(15.2))


def _cleanup(paths):
    for path in paths:
        try:
            if path and os.path.isfile(path):
                os.remove(path)
        except OSError:
            pass


def _format_number(value):
    if isinstance(value, float) and value != int(value):
        return ("%.1f" % value).rstrip("0").rstrip(".")
    try:
        return str(int(value))
    except (TypeError, ValueError):
        return str(value)


def _append_others(lead, intro, parts):
    if not parts:
        return lead
    return "%s%s%s。" % (lead, intro, "；".join(parts))


def _hour_detail_text(hour_counts, peak_hour, peak_count):
    ranked = sorted(
        [
            (hour, int(count))
            for hour, count in enumerate(hour_counts or [])
            if int(count) > 0
        ],
        key=lambda item: (-item[1], item[0]),
    )
    lead = "发生起数最多的小时为 %d 时，该小时发生事故 %d 起。" % (
        int(peak_hour or 0),
        int(peak_count or 0),
    )
    others = ranked[1:]
    parts = ["%d时%d起" % (hour, count) for hour, count in others]
    return _append_others(lead, "其他时段的情况分别为", parts)


def _hotspot_detail_text(hotspot_rows):
    top = hotspot_rows[0]
    lead = "事故最多的路段为%s，共 %s 起。" % (
        top["location"],
        _format_number(top["count"]),
    )
    parts = [
        "%s共%s起" % (item["location"], _format_number(item["count"]))
        for item in hotspot_rows[1:]
    ]
    return _append_others(lead, "其他路段的情况分别为", parts)


def _category_detail_text(block, lead, intro):
    items = list(block.get("items") or [])
    nonempty = int(block.get("nonempty") or 0)
    parts = []
    for label, count in items[1:]:
        pct = (100.0 * count / nonempty) if nonempty else 0.0
        parts.append("%s占 %.1f%%，共 %s 起" % (label, pct, _format_number(count)))
    return _append_others(lead, intro, parts)


def write_report_docx(dest_path, data):
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt

    document = Document()
    section = document.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.left_margin = Mm(25)
    section.right_margin = Mm(25)
    section.top_margin = Mm(25)
    section.bottom_margin = Mm(25)

    style = document.styles["Normal"]
    style.font.name = "仿宋"
    style.font.size = Pt(16)
    r_pr = style.element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn("w:eastAsia"), "仿宋")

    temp_images = []
    try:
        _add_title(document, "事故分析报告")

        _add_h1(document, "一、事故整体情况")
        overview = (
            "共发生事故 %d 起，一般程序 %d 起，简易程序 %d 起，造成 %s 人死亡，%s 人受伤。"
            % (
                data["total"],
                data["general"],
                data["simple"],
                _format_number(data["deaths"]),
                _format_number(data["injuries"]),
            )
        )
        _add_body(document, overview)

        _add_h1(document, "二、事故道路分布情况")
        road_rows = data.get("road_rows") or []
        if not road_rows:
            _add_body(document, NO_DATA)
        else:
            _add_table(
                document,
                ["道路", "事故数", "死亡人数", "受伤人数"],
                [
                    [
                        item["name"],
                        _format_number(item["count"]),
                        _format_number(item["deaths"]),
                        _format_number(item["injuries"]),
                    ]
                    for item in road_rows
                ],
            )

        _add_h1(document, "三、事故管辖分布情况")
        jurisdiction_rows = data.get("jurisdiction_rows") or []
        if not jurisdiction_rows:
            _add_body(document, NO_DATA)
        else:
            _add_table(
                document,
                ["管辖", "事故数", "死亡人数", "受伤人数"],
                [
                    [
                        item["name"],
                        _format_number(item["count"]),
                        _format_number(item["deaths"]),
                        _format_number(item["injuries"]),
                    ]
                    for item in jurisdiction_rows
                ],
            )

        _add_h1(document, "四、详细分析")

        _add_h2(document, "（一）事故时间分布")
        if not data.get("time_ok"):
            _add_body(
                document,
                "未能解析「时间」列中的日期时间。请确认该列为日期、时间或常见文本格式（例如 2024-01-01 13:20）。",
            )
        else:
            image = hour_line_chart(data.get("hour_counts") or [])
            if image:
                temp_images.append(image)
                _add_picture(document, image)
            _add_body(
                document,
                _hour_detail_text(
                    data.get("hour_counts") or [],
                    data.get("peak_hour") or 0,
                    data.get("peak_count") or 0,
                ),
            )

        _add_h2(document, "（二）事故多发路段")
        hotspot_rows = data.get("hotspot_rows") or []
        if not hotspot_rows:
            _add_body(document, NO_DATA)
        else:
            _add_table(
                document,
                ["路段位置", "事故数"],
                [
                    [item["location"], _format_number(item["count"])]
                    for item in hotspot_rows
                ],
            )
            _add_body(document, _hotspot_detail_text(hotspot_rows))

        cause = (data.get("optional") or {}).get("cause") or {}
        _add_h2(document, "（三）事故原因")
        if not cause.get("mapped"):
            _add_body(document, CANNOT_ANALYZE)
        elif not cause.get("nonempty") or not cause.get("items"):
            _add_body(document, NO_DATA)
        else:
            image = pie_chart(cause["items"], "事故原因")
            if image:
                temp_images.append(image)
                _add_picture(document, image)
            _add_body(
                document,
                _category_detail_text(
                    cause,
                    "占比最高的原因为%s，占 %.1f%%，共 %s 起。"
                    % (
                        cause.get("top_label") or "",
                        float(cause.get("top_pct") or 0),
                        _format_number(cause.get("top_count") or 0),
                    ),
                    "其他原因分别为",
                ),
            )

        weather = (data.get("optional") or {}).get("weather") or {}
        _add_h2(document, "（四）天气分析")
        if not weather.get("mapped") or not weather.get("nonempty") or not weather.get("items"):
            _add_body(document, NO_DATA)
        else:
            image = bar_chart(weather["items"], "天气分析")
            if image:
                temp_images.append(image)
                _add_picture(document, image)
            _add_body(
                document,
                _category_detail_text(
                    weather,
                    "最主要天气为%s，占 %.1f%%，共 %s 起。"
                    % (
                        weather.get("top_label") or "",
                        float(weather.get("top_pct") or 0),
                        _format_number(weather.get("top_count") or 0),
                    ),
                    "其他天气分别为",
                ),
            )

        form = (data.get("optional") or {}).get("form") or {}
        _add_h2(document, "（五）事故形态分析")
        if not form.get("mapped") or not form.get("nonempty") or not form.get("items"):
            _add_body(document, NO_DATA)
        else:
            image = pie_chart(form["items"], "事故形态")
            if image:
                temp_images.append(image)
                _add_picture(document, image)
            _add_body(
                document,
                _category_detail_text(
                    form,
                    "主要形态为%s，占 %.1f%%，共 %s 起。"
                    % (
                        form.get("top_label") or "",
                        float(form.get("top_pct") or 0),
                        _format_number(form.get("top_count") or 0),
                    ),
                    "其他形态分别为",
                ),
            )

        document.save(dest_path)
    finally:
        _cleanup(temp_images)


def generate_report(headers, rows, mapping, src_path, progress_cb=None, should_cancel=None):
    _raise_if_cancelled(should_cancel)
    if progress_cb:
        progress_cb(12, "正在统计事故数据")
    data = build_report_data(headers, rows, mapping)
    _raise_if_cancelled(should_cancel)
    if progress_cb:
        progress_cb(45, "正在生成图表与报告")
    dest = output_path_for(src_path)
    write_report_docx(dest, data)
    _raise_if_cancelled(should_cancel)
    if progress_cb:
        progress_cb(100, "报告生成完成")
    return dest
